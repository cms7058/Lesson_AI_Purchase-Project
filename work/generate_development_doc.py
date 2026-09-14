from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "制造业采购AI智能体_技术开发设计说明书_V0.1.docx"
FONT = "PingFang SC"
BLUE = RGBColor(46, 116, 181)
DARK = RGBColor(31, 77, 120)
INK = RGBColor(16, 24, 40)
MUTED = RGBColor(102, 112, 133)
LIGHT = "E8EEF5"
PALE = "F4F6F9"


def set_font(run, size=11, bold=False, color=INK, name=FONT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(doc, headers, rows, widths, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade(cell, LIGHT)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(header), size=font_size, bold=True, color=DARK)
    repeat_header(table.rows[0])
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[index].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            set_font(p.add_run(str(value)), size=font_size)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_para(doc, text, size=11, bold=False, color=INK, after=6, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if align is not None:
        p.alignment = align
    set_font(p.add_run(text), size=size, bold=bold, color=color)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Inches(0.375 if level == 0 else 0.625)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    set_font(p.add_run(text), size=10.5)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p.add_run(text)


def add_callout(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    shade(cell, PALE)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run(title + "  "), size=10.5, bold=True, color=DARK)
    set_font(p.add_run(text), size=10.5)
    set_table_geometry(table, [9360])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    settings = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK, 10, 5),
    }
    for name, (size, color, before, after) in settings.items():
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, fld_end])
    set_font(run, size=9, color=MUTED)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    configure_styles(doc)

    header = section.header.paragraphs[0]
    header.paragraph_format.space_after = Pt(0)
    set_font(header.add_run("PEBS COPILOT PURCHASE  |  技术开发设计"), size=9, color=MUTED)
    add_page_number(section.footer.paragraphs[0])

    add_para(doc, "技术开发设计说明书", size=26, bold=True, color=DARK, after=5)
    add_para(doc, "制造业采购AI智能体（PEBS Copilot Purchase）", size=15, bold=True, color=BLUE, after=18)
    add_table(
        doc,
        ["项目", "内容"],
        [
            ("文档版本", "V0.1 - 开发基线"),
            ("编制日期", "2026年9月3日"),
            ("前端约束", "Vue 2.7.16；依赖锁定并保留Vue 3迁移路径"),
            ("后端框架", "Python FastAPI"),
            ("建设范围", "采购分析、执行、自动化、预测、寻源与多工厂优化"),
            ("当前里程碑", "M0.1：工程骨架与采购订单纵向切片"),
        ],
        [1750, 7610],
        font_size=10,
    )
    add_callout(doc, "开发原则", "大模型负责理解、编排和解释；价格、TCO、回归、预测和路径优化由可验证的程序或算法执行。定标、签署、订单发送等影响性动作必须经过权限校验和人工确认。")
    doc.add_page_break()

    add_heading(doc, "1. 文档目的与开发范围")
    add_para(doc, "本文将已确认的系统需求转换为可实施的技术架构、领域模块、数据实体、接口边界、开发里程碑和质量标准，作为代码设计、接口联调、测试验收和变更控制的开发基线。")
    add_heading(doc, "1.1 完整模块范围", 2)
    modules = [
        ("M01", "AI采购助手", "查数、分析、模型修正、任务优化、报告和受控执行"),
        ("M02", "数据与API接入", "文件导入、系统连接器、映射、全量/增量同步和审计"),
        ("M03", "报价与比价", "OCR、标准化、自动比价、TCO和招标评价"),
        ("M04", "采购订单", "创建、审批、发送、确认、变更、交付、收货和关闭"),
        ("M05", "模板中心", "订单、报价单和合同模板、变量映射、版本和生成"),
        ("M06", "自动化工作流", "报价、订单、合同签署、邮件、审批和异常恢复"),
        ("M07", "需求预测", "多模型预测、采购建议、预警和情景模拟"),
        ("M08", "智能寻源", "外部候选发现、能力匹配、风险初筛和准入"),
        ("M09", "供应商管理", "360档案、绩效、风险、关系和价值"),
        ("M10", "多工厂路径优化", "供货分配、路线批次、约束和中断重算"),
        ("M11", "合同管理", "生成、审查、审批、电子签署和履约提醒"),
        ("M12", "报告通知与审计", "驾驶舱、报告、邮件、站内消息和全链路留痕"),
    ]
    add_table(doc, ["编号", "模块", "开发边界"], modules, [900, 2200, 6260])

    add_heading(doc, "2. 总体技术架构")
    add_para(doc, "系统采用前后端分离、模块化单体起步、算法服务解耦的架构。业务量和团队规模增长后，可按数据接入、文档处理、AI编排、预测和优化服务逐步拆分。")
    add_callout(doc, "运行链路", "Vue管理端 → FastAPI API层 → 领域服务/AI编排/工作流 → PostgreSQL、Redis、对象存储 → 大模型、OCR、ERP/SRM/MES/WMS/QMS/OA及电子签署平台。")
    add_heading(doc, "2.1 技术选型", 2)
    stack = [
        ("前端", "Vue 2.7.16、Vue Router、Vuex、Axios、Element UI、ECharts", "采购工作台、流程交互和可视化"),
        ("后端", "FastAPI、Pydantic、SQLAlchemy、Alembic", "API、校验、领域服务和数据库迁移"),
        ("数据", "PostgreSQL、pgvector、Redis、MinIO", "业务、知识、缓存、会话和文件"),
        ("任务", "Celery + Redis", "OCR、同步、报告和模型训练"),
        ("流程", "Camunda 8 + BPMN/DMN；前端bpmn-js", "用户自定义采购自动化"),
        ("AI", "模型适配层 + LangGraph（仅复杂长流程）", "模型可替换、工具调用、人工中断和恢复"),
        ("算法", "Pandas、statsmodels、scikit-learn、OR-Tools", "回归、预测、分配和路线优化"),
        ("部署", "Docker Compose；生产按需Kubernetes", "环境一致性、扩容和运维"),
    ]
    add_table(doc, ["层级", "技术", "作用"], stack, [1200, 4100, 4060], font_size=8.8)

    add_heading(doc, "3. 工程结构与模块边界")
    add_para(doc, "仓库采用apps/api与apps/web双应用结构。后端按api、core、domain、services和后续repositories分层；前端按views、components、router、store和api组织。")
    add_table(doc, ["目录", "职责"], [
        ("apps/api/app/api", "HTTP接口、请求响应模型和权限依赖"),
        ("apps/api/app/domain", "订单、报价、合同、供应商等领域模型"),
        ("apps/api/app/services", "业务规则、任务编排和第三方适配"),
        ("apps/api/app/repositories", "数据库持久化接口；M0.2引入"),
        ("apps/web/src/views", "业务页面和工作台"),
        ("apps/web/src/api", "统一API客户端、错误和认证处理"),
        ("docs", "架构决策、接口约定和开发记录"),
    ], [2800, 6560])

    add_heading(doc, "4. 核心领域数据设计")
    entities = [
        ("组织权限", "organization、factory、department、user、role、permission"),
        ("主数据", "material、category、unit、currency、warehouse、supplier"),
        ("询报价", "rfq、rfq_supplier、quotation、quotation_line、comparison"),
        ("合同", "contract、contract_version、clause、review、signature_event"),
        ("订单", "purchase_order、order_line、order_change、delivery、receipt、inspection"),
        ("模板", "document_template、template_version、template_field、generation_record"),
        ("工作流", "workflow_definition、workflow_version、workflow_instance、task、event"),
        ("预测", "forecast_dataset、forecast_run、forecast_result、purchase_suggestion"),
        ("寻源", "sourcing_task、candidate_supplier、evidence、qualification"),
        ("优化", "network_node、route、capacity、optimization_run、allocation_result"),
        ("AI审计", "conversation、message、tool_call、model_run、evidence、feedback"),
    ]
    add_table(doc, ["数据域", "核心实体"], entities, [2000, 7360], font_size=9)
    add_bullet(doc, "所有业务数据必须包含组织范围、创建人、创建时间、更新时间和版本字段。")
    add_bullet(doc, "报价、合同、订单、模型结果及工作流实例采用逻辑删除或不可变版本，不直接物理覆盖。")
    add_bullet(doc, "外部数据保存源系统、源记录键、同步批次和映射版本，保证幂等和血缘追踪。")

    add_heading(doc, "5. 采购订单纵向切片")
    add_heading(doc, "5.1 订单状态机", 2)
    add_para(doc, "草稿 → 待审批 → 已审批 → 待发送 → 已发送 → 供应商确认 → 部分交货 → 已交货 → 已收货/质检 → 已完成。驳回、暂停、取消和异常为受控分支。")
    add_heading(doc, "5.2 首批接口", 2)
    add_table(doc, ["方法", "路径", "用途"], [
        ("GET", "/api/v1/health", "服务健康检查"),
        ("GET", "/api/v1/modules", "返回完整模块目录及状态"),
        ("GET", "/api/v1/orders", "查询采购订单"),
        ("POST", "/api/v1/orders", "创建订单并计算含税金额"),
        ("POST", "/api/v1/assistant/chat", "提交AI采购助手请求"),
    ], [1000, 3300, 5060])
    add_callout(doc, "当前实现边界", "M0.1订单数据使用进程内仓储验证接口和页面闭环；M0.2替换为SQLAlchemy/PostgreSQL持久化，并增加认证、权限、审批和状态迁移。")

    add_heading(doc, "6. 模板与文档生成服务")
    add_bullet(doc, "支持Word、Excel和可填写PDF模板；扫描PDF仅作为识别与参考材料。")
    add_bullet(doc, "模板变量、明细循环、条件片段、金额大小写、签章位置和多语言格式可配置。")
    add_bullet(doc, "AI辅助识别字段，用户确认映射；模板发布需版本、生效范围和审批。")
    add_bullet(doc, "每次生成保存模板版本、数据快照、生成文件、操作者和业务对象关联。")

    add_heading(doc, "7. 自动化工作流设计")
    add_para(doc, "业务用户通过BPMN设计器配置触发器、审批、条件、并行、AI任务、系统API、文档生成、邮件、电子签署、等待、重试和人工接管节点。")
    add_table(doc, ["流程", "标准链路"], [
        ("报价自动化", "需求确认→生成询价→审批→邮件发送→报价接收→识别→比价→定标确认"),
        ("订单自动化", "中标/合同→生成订单→价格校验→审批→发送→供应商确认→履约跟踪"),
        ("合同签署", "生成合同→AI审查→业务/法务审批→电子签署→回调校验→归档提醒"),
    ], [1800, 7560])
    add_bullet(doc, "流程定义按版本发布，运行中的实例继续使用启动时版本。")
    add_bullet(doc, "外部发送和签署必须支持白名单、预览确认、幂等键、失败重试和人工补偿。")

    add_heading(doc, "8. AI助手与模型治理")
    add_bullet(doc, "建立模型供应商适配层，业务代码不得直接绑定单一模型SDK。")
    add_bullet(doc, "工具调用采用白名单和结构化参数，先做权限检查再执行。")
    add_bullet(doc, "知识问答返回证据来源、数据时间和业务范围；无证据时明确提示。")
    add_bullet(doc, "模型修正进入反馈、评测、审批、发布和回滚闭环。")
    add_bullet(doc, "未配置模型时系统返回配置状态，不生成伪造的业务结论。")

    add_heading(doc, "9. 预测、寻源与路径优化")
    add_heading(doc, "9.1 需求预测", 2)
    add_para(doc, "统一历史领用、订单、生产计划、库存、在途和供应提前期，比较移动平均、指数平滑、季节性及回归模型，通过误差和回测选择方案，并输出采购建议与情景模拟。")
    add_heading(doc, "9.2 外部智能寻源", 2)
    add_para(doc, "寻源服务将自然语言条件结构化，从企业库和经批准外部数据源获取候选供应商，完成实体去重、证据归档、能力匹配、风险初筛、长短名单和准入自动化。")
    add_heading(doc, "9.3 多工厂优化", 2)
    add_para(doc, "OR-Tools模型同时处理供应商产能、工厂需求、MOQ、交期、仓储、路线和供应占比约束，以采购、运输、延期、风险和碳排的加权成本为目标，输出供货量、路线、批次及备选方案。")

    add_heading(doc, "10. 安全、权限与审计")
    security = [
        ("身份", "企业SSO/OIDC；开发环境可使用本地认证"),
        ("授权", "集团、公司、工厂、采购组织、品类和字段级权限"),
        ("密钥", "API、模型、邮件和签署凭据进入密钥管理，不写入代码或日志"),
        ("数据", "TLS传输、敏感字段加密、导出水印、保留与删除策略"),
        ("AI", "提示注入防护、工具白名单、数据脱敏和模型路由"),
        ("审计", "查询、导入、修改、生成、审批、发送、签署和模型调用全记录"),
    ]
    add_table(doc, ["控制域", "要求"], security, [1800, 7560])

    add_heading(doc, "11. 开发里程碑")
    milestones = [
        ("M0.1", "工程启动", "前后端骨架、模块目录、订单创建查询、AI安全占位、基础设施定义", "已启动"),
        ("M0.2", "平台底座", "PostgreSQL、迁移、认证权限、对象存储、任务队列、审计", "下一步"),
        ("M1", "采购执行闭环", "数据接入、报价比价、TCO回归、模板、合同、订单、邮件和工作流", "计划"),
        ("M2", "预测与供应商智能", "需求预测、采购建议、外部寻源、准入、风险和绩效", "计划"),
        ("M3", "多工厂优化", "供货分配、路径批次、情景模拟、动态订单调整", "计划"),
    ]
    add_table(doc, ["阶段", "主题", "交付内容", "状态"], milestones, [900, 1700, 5560, 1200], font_size=8.8)

    add_heading(doc, "12. 测试与验收基线")
    add_bullet(doc, "后端：单元测试、API集成测试、权限测试、幂等与失败恢复测试。")
    add_bullet(doc, "前端：构建检查、核心页面组件测试和关键流程端到端测试。")
    add_bullet(doc, "算法：固定数据集、回测指标、模型版本、基线对比和可复现结果。")
    add_bullet(doc, "文档/OCR：字段准确率、来源定位、人工校正和异常样本集。")
    add_bullet(doc, "安全：依赖扫描、密钥检查、越权、上传文件和提示注入测试。")
    add_bullet(doc, "性能：普通API、AI首段响应、批量同步和长任务容量分别验收。")

    doc.add_page_break()
    add_heading(doc, "13. 待联调确认")
    confirmations = [
        ("首个业务系统", "ERP/SRM名称、接口文档、测试环境、鉴权和增量规则"),
        ("模板样本", "订单、报价单、合同各3至5份典型可编辑模板"),
        ("邮件", "SMTP、Microsoft 365或企业邮件网关及测试账户"),
        ("电子签署", "e签宝、法大大、上上签或客户指定平台"),
        ("模型", "公有云、私有化或混合路由，及敏感数据策略"),
        ("成本口径", "TCO项目、实际采购成本定义和回归训练样本"),
        ("优化数据", "供应商产能、工厂需求、路线、运费、时效和约束"),
    ]
    add_table(doc, ["主题", "需要客户提供或确认"], confirmations, [2200, 7160])

    doc.core_properties.title = "制造业采购AI智能体技术开发设计说明书"
    doc.core_properties.subject = "PEBS Copilot Purchase开发基线"
    doc.core_properties.author = "PEBS项目组"
    doc.core_properties.keywords = "采购AI, FastAPI, Vue2, 工作流, 订单, 需求预测, 智能寻源, 路径优化"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
