from pathlib import Path
from datetime import date

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path("/Users/mingyue/PEBS_Copilot_Purchase")
OUT = ROOT / "制造业采购AI智能体_系统需求规格说明书_V1.2.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "203040"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
WHITE = "FFFFFF"
RED = "9B1C1C"
GOLD = "7A5A00"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent=120):
    total = sum(widths_dxa)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size=11, bold=False, color=INK, name="Arial Unicode MS"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, 9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    tail = paragraph.add_run(" 页")
    set_run_font(tail, 9, color=MUTED)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_para(doc, text="", bold_prefix=None, after=6, align=None, color=INK, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    if align is not None:
        p.alignment = align
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=size, bold=True, color=color)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, size=size, color=color)
    else:
        r = p.add_run(text)
        set_run_font(r, size=size, color=color)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.10
    run = p.add_run(text)
    set_run_font(run)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.10
    run = p.add_run(text)
    set_run_font(run)
    return p


def add_requirement(doc, rid, title, description, priority="必须"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.10
    keep_with_next(p)
    tag = p.add_run(f"{rid}｜{priority}｜{title}：")
    set_run_font(tag, bold=True, color=DARK_BLUE)
    body = p.add_run(description)
    set_run_font(body)
    return p


def add_heading(doc, text, level):
    p = doc.add_heading(text, level=level)
    keep_with_next(p)
    return p


def add_table(doc, headers, rows, widths, header_fill=LIGHT_BLUE, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, value in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, header_fill)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(value)
        set_run_font(r, size=font_size, bold=True, color=DARK_BLUE)
    for row_data in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_data):
            cell = cells[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_run_font(r, size=font_size)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial Unicode MS"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial Unicode MS")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial Unicode MS")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Arial Unicode MS"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial Unicode MS")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial Unicode MS")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ["List Bullet", "List Bullet 2", "List Number"]:
        style = styles[name]
        style.font.name = "Arial Unicode MS"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial Unicode MS")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial Unicode MS")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.10


def build_document():
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(0.88)
    section.right_margin = Inches(0.88)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("制造业采购AI智能体｜系统需求规格说明书")
    set_run_font(hr, 8.5, color=MUTED)
    footer = section.footer
    add_page_number(footer.paragraphs[0])

    # First-page masthead
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run("系统需求规格说明书")
    set_run_font(r, size=25, bold=True, color=DARK_BLUE)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(18)
    r2 = p2.add_run("制造业采购AI智能体（PEBS Copilot Purchase）")
    set_run_font(r2, size=15, bold=True, color=BLUE)

    metadata = [
        ("文档版本", "V1.2"),
        ("需求状态", "需求评审稿"),
        ("编制日期", "2026年9月3日"),
        ("适用对象", "项目发起人、客户代表、产品经理、架构师、开发与测试团队"),
        ("需求基线", "V1.1需求规格说明书，并补充系统API数据接入能力"),
        ("本次变更", "新增系统API数据导入、全量/增量同步、字段映射、鉴权、重试、幂等与同步审计要求"),
    ]
    add_table(doc, ["项目", "内容"], metadata, [1600, 7760], font_size=10)

    p3 = doc.add_paragraph()
    p3.paragraph_format.space_before = Pt(12)
    p3.paragraph_format.space_after = Pt(8)
    r3 = p3.add_run("文档目的")
    set_run_font(r3, 12, bold=True, color=DARK_BLUE)
    add_para(
        doc,
        "本说明书定义制造业采购AI智能体的业务目标、用户角色、功能范围、数据要求、接口要求、非功能要求及验收边界，作为需求评审、产品设计、开发估算、测试验收和后续变更控制的依据。",
    )

    p4 = doc.add_paragraph()
    p4.paragraph_format.space_before = Pt(8)
    p4.paragraph_format.space_after = Pt(4)
    r4 = p4.add_run("本轮交付边界")
    set_run_font(r4, 12, bold=True, color=DARK_BLUE)
    add_bullet(doc, "本文件为系统需求文档，不包含数据库表结构、API字段定义、代码目录、部署脚本等开发设计内容。")
    add_bullet(doc, "后端拟采用 FastAPI，前端按客户要求采用 Vue 2.0；最终技术方案须在客户确认后写入开发文档。")
    add_bullet(doc, "AI输出属于辅助决策，招标定标、供应商准入、合同审批等关键动作必须由有权限的人员确认。")

    doc.add_page_break()

    add_heading(doc, "1. 项目概述", 1)
    add_heading(doc, "1.1 建设背景", 2)
    add_para(
        doc,
        "制造业采购工作涉及需求计划、市场分析、供应商寻源、询报价、招标比价、成本核算、合同审查、交付与质量跟踪等大量跨系统、跨文档任务。现有工作普遍存在数据格式不统一、人工整理时间长、价格与风险判断依赖个人经验、历史知识难复用、采购决策依据不完整等问题。",
    )
    add_para(
        doc,
        "本项目拟建设面向制造业采购场景的AI智能体，将大语言模型、企业知识库、文档识别、统计分析、优化算法与采购业务流程结合，形成可查询、可计算、可追溯、可审批的采购辅助决策平台。",
    )

    add_heading(doc, "1.2 建设目标", 2)
    goals = [
        "降低报价、合同、历史订单和供应商资料的人工整理成本。",
        "将市场、产业链、价格、质量、交期与服务数据转化为可解释的采购建议。",
        "形成需求预测、TCO分析、招标比价、供应商评价和风险预警的标准化能力。",
        "打通采购知识与业务数据，沉淀企业采购经验并支持多人协同复用。",
        "保留完整数据来源、计算过程和审批痕迹，满足审计、合规和责任追溯要求。",
    ]
    for item in goals:
        add_bullet(doc, item)

    add_heading(doc, "1.3 系统定位", 2)
    add_para(
        doc,
        "系统定位为“采购决策助手与业务协同平台”，覆盖市场洞察、需求预测、智能寻源、询报价、招标比价、采购成本、供应商全生命周期、合同审查和履约监控。系统不替代ERP、SRM、QMS等业务系统的法定记账和审批职责，也不允许大模型绕过企业授权自动完成中标、签约或付款。",
    )

    add_heading(doc, "1.4 建设范围与阶段", 2)
    phase_rows = [
        ("一期", "基础可用", "AI问答、企业知识库、文件/API数据导入、报价整理、智能比价、历史价格分析、供应商绩效、合同审查、报告导出"),
        ("二期", "决策增强", "需求预测、TCO模型、多元回归招标评价、供应商风险预警、智能寻源、准入管理、外部数据接入"),
        ("三期", "优化协同", "产业链知识图谱、多工厂供货分配、运输路径优化、供应中断模拟、动态采购策略"),
    ]
    add_table(doc, ["阶段", "目标", "范围"], phase_rows, [900, 1500, 6960])

    add_heading(doc, "2. 术语与业务原则", 1)
    term_rows = [
        ("AI智能体", "能理解采购任务、调用知识、数据、模型或业务接口并生成结果的协作式软件能力。"),
        ("RAG", "检索增强生成。先从授权知识库检索证据，再由大模型生成答案。"),
        ("TCO", "总拥有成本。除采购单价外，还包括物流、资金、质量、延期、服务等显性和隐性成本。"),
        ("成品率", "合格成品数量占投入或交付数量的比例，具体口径由企业确认。"),
        ("返修率", "需要返工或返修的数量占检验或交付数量的比例。"),
        ("多元线性回归", "使用多个自变量共同解释或预测目标变量的统计模型。本项目用于估计实际采购成本。"),
        ("人工复核", "AI生成建议后，由具备相应权限的人员审阅、修改并确认。"),
    ]
    add_table(doc, ["术语", "定义"], term_rows, [1800, 7560])

    add_heading(doc, "2.1 核心业务原则", 2)
    principles = [
        "有据可查：结论应展示引用的文件、数据时间、计算口径和模型版本。",
        "数据优先：结构化计算、统计模型和优化算法不得由大语言模型凭空代算。",
        "人在回路：关键业务决策必须保留人工确认和审批。",
        "权限最小化：用户、智能体及外部接口仅可访问完成任务所需的最小数据范围。",
        "结果可复现：同一数据、同一规则和同一模型版本应能够复现关键评分与计算结果。",
    ]
    for item in principles:
        add_bullet(doc, item)

    add_heading(doc, "3. 用户角色与权限", 1)
    roles = [
        ("采购专员", "发起问答、上传报价、创建比价、维护寻源任务、查看本人授权的供应商和合同。"),
        ("采购经理", "审核采购分析、调整评分权重、批准推荐结论、查看团队与供应商绩效。"),
        ("需求/计划人员", "维护需求计划、查看预测结果、确认异常需求和建议采购量。"),
        ("质量人员", "维护检验、不良、返修和改善数据，参与供应商质量评价。"),
        ("法务/合同人员", "维护合同模板与条款库，复核合同风险和修改建议。"),
        ("供应链管理层", "查看采购驾驶舱、重大风险、降本收益和跨工厂供应方案。"),
        ("系统管理员", "维护组织、用户、角色、模型、字典、接口、审计和系统参数。"),
        ("审计人员", "只读查看审批记录、模型依据、操作日志和异常合规报告。"),
    ]
    add_table(doc, ["角色", "主要权限"], roles, [1900, 7460])

    add_requirement(doc, "AUTH-001", "组织与角色隔离", "系统应支持按集团、公司、工厂、采购组织、品类和岗位配置数据权限。")
    add_requirement(doc, "AUTH-002", "敏感字段权限", "报价、成本、银行账户、联系方式和合同内容应支持字段级查看与脱敏。")
    add_requirement(doc, "AUTH-003", "关键动作控制", "招标推荐确认、供应商准入、合同定稿、模型发布和数据导出应支持单独授权。")

    add_heading(doc, "4. 总体业务流程", 1)
    flow = [
        "采购人员以自然语言、表单、上传文件或选择已配置系统API数据源的方式发起任务。",
        "系统识别任务意图、权限范围及所需的数据或工具。",
        "系统从知识库、文件导入、系统API同步数据和经授权的外部数据源检索证据。",
        "系统调用文档解析、统计模型、价格计算、风险规则或优化算法执行任务。",
        "系统生成结构化结果、依据、风险提示和建议下一步。",
        "有权限的业务人员复核、修改、确认或驳回结果。",
        "系统保存输入、证据、计算结果、模型版本和审批日志，并按权限回写业务系统。",
    ]
    for item in flow:
        add_number(doc, item)

    add_heading(doc, "5. 功能需求", 1)

    add_heading(doc, "5.1 对话式AI采购助手与任务编排", 2)
    add_requirement(doc, "FR-AI-001", "采购自然语言问答", "用户可围绕物料、价格、供应商、合同、市场、质量、交期和采购制度进行中文问答及连续追问。")
    add_requirement(doc, "FR-AI-002", "采购提示词模板", "系统应提供数据整理、市场分析、产业链分析、报价分析、谈判准备、风险识别和汇报生成等可复用模板。")
    add_requirement(doc, "FR-AI-003", "任务识别与拆解", "系统应将复杂问题拆解为检索、解析、计算、比较、生成和审批等步骤，并向用户展示任务状态。")
    add_requirement(doc, "FR-AI-004", "证据引用", "回答应关联来源文件、业务记录、数据日期和关键片段；无法获得可靠依据时必须明确提示。")
    add_requirement(doc, "FR-AI-005", "会话与任务记录", "用户应可保存、命名、继续、复制和归档会话，并将有效结论转存为采购知识。")
    add_requirement(doc, "FR-AI-006", "对话式报告生成", "用户可通过对话指定主题、数据范围、时间、模板和受众，将问答及分析结果生成采购简报、市场报告、比价说明、谈判提纲或管理汇报。")
    add_requirement(doc, "FR-AI-007", "人工反馈", "用户可对答案标记正确、错误、需补充或不适用，反馈应进入可审计的优化闭环。", "应该")
    add_requirement(doc, "FR-AI-008", "统一AI助手入口", "系统应在主工作台提供常驻AI助手，用户可通过连续对话、快捷指令、上传附件或引用当前页面数据发起任务。")
    add_requirement(doc, "FR-AI-009", "对话式数据查找", "用户可用自然语言查找物料、供应商、报价、订单、合同、交付、质量、库存和市场数据；AI助手应按用户权限返回结构化结果、筛选条件、数据时间及来源。")
    add_requirement(doc, "FR-AI-010", "上下文与范围控制", "AI助手应理解当前会话、页面对象、附件和用户选择的数据范围，支持追问、改写条件和恢复上下文；切换组织或敏感数据范围时应重新校验权限。")
    add_requirement(doc, "FR-AI-011", "模型修正闭环", "用户可通过对话提交错误样例、补充正确答案、调整规则或权重、发起数据重算或模型重训申请；系统应保存修改原因、影响范围、验证结果和版本，生产模型发布或回滚必须经授权审批。")
    add_requirement(doc, "FR-AI-012", "任务优化", "AI助手应分析采购任务的目标、约束、依赖、优先级、执行步骤和异常瓶颈，提出合并、拆分、排序、补充数据、替代工具或重新执行建议，并在用户确认后执行可授权的优化动作。")
    add_requirement(doc, "FR-AI-013", "报告编辑与版本", "用户应可通过对话要求增加、删除、改写、重新计算或更换报告模板；系统应保留报告版本、数据口径、引用来源、生成时间和人工修改记录。")
    add_requirement(doc, "FR-AI-014", "执行预览与确认", "AI助手调用数据写入、模型调整、任务重跑、批量导出或对外发送等有影响的工具前，应展示执行对象、范围和预期影响，并按权限要求获得用户确认。")
    add_requirement(doc, "FR-AI-015", "任务进度与异常恢复", "长时间运行的数据查找、模型计算、任务优化和报告生成应进入任务中心，展示步骤、进度和错误；失败后支持安全重试、从中断点继续或转人工处理。")

    add_heading(doc, "5.2 企业知识库与采购数据整理", 2)
    add_requirement(doc, "FR-DATA-001", "多格式导入", "系统应支持Excel、CSV、Word、PDF、图片、扫描件及压缩包批量导入，并保留原文件。")
    add_requirement(doc, "FR-DATA-002", "文档智能识别", "系统应提取物料、规格、品牌、数量、单位、币种、税率、价格、交期、付款、有效期和供应商信息。")
    add_requirement(doc, "FR-DATA-003", "主数据标准化", "系统应统一供应商名称、物料编码、计量单位、币种、税率和日期格式，并保留标准化前后的映射。")
    add_requirement(doc, "FR-DATA-004", "数据质量检查", "系统应检测缺失、重复、冲突、异常值和识别置信度不足，并允许人工校正。")
    add_requirement(doc, "FR-DATA-005", "知识库管理", "管理员应可按制度、合同、供应商、物料、行业和项目维护知识库，配置可见范围和生效日期。")
    add_requirement(doc, "FR-DATA-006", "版本与失效控制", "文件更新后应保留历史版本；过期制度、合同模板或供应商资质不得作为当前有效依据。")
    add_requirement(doc, "FR-DATA-007", "数据血缘", "分析结果应可追溯至原始文件、源系统记录、字段映射和人工修改记录。")
    add_requirement(doc, "FR-DATA-008", "系统API数据导入", "系统应支持通过客户授权的ERP、SRM、MES、WMS、QMS、OA及其他业务系统API导入主数据、交易、库存、质量、交付和审批数据，并将API导入与文件导入纳入统一数据处理流程。")
    add_requirement(doc, "FR-DATA-009", "同步策略", "系统应支持全量、增量、手动、定时及事件触发同步；同步任务可配置时间范围、分页、批次大小、变更时间戳或水位标识，并避免遗漏和重复读取。")
    add_requirement(doc, "FR-DATA-010", "字段映射与转换", "管理员应可配置接口地址、请求参数、分页规则、源字段至目标字段映射、数据类型、单位、编码转换、默认值和校验规则；映射配置应具备版本、生效与回滚能力。")
    add_requirement(doc, "FR-DATA-011", "接口鉴权与密钥安全", "数据导入接口应支持API Key、OAuth2、客户端证书或签名等客户认可的鉴权方式；凭据应加密保存或进入专用密钥管理，不得在页面、日志或错误信息中明文展示。")
    add_requirement(doc, "FR-DATA-012", "同步监控与异常恢复", "系统应展示同步任务进度、读取数、成功数、失败数、跳过数及错误明细；发生超时、限流、网络中断或数据错误时，应支持退避重试、断点续传、人工重跑和失败记录隔离。")
    add_requirement(doc, "FR-DATA-013", "API数据血缘与幂等", "系统应记录源系统、接口端点、请求时间、同步批次、源记录主键和同步版本；同一业务记录重复到达时不得产生重复数据，并应按配置安全更新或拒绝冲突。")

    add_heading(doc, "5.3 市场环境与产业链分析", 2)
    add_requirement(doc, "FR-MKT-001", "市场专题分析", "系统应支持按行业、物料、区域和时间范围生成供需、价格、竞争、政策、汇率、关税和运输环境分析。")
    add_requirement(doc, "FR-MKT-002", "航空维修行业支持", "系统应支持航空维修场景下的原厂件、PMA件、适航认证、航材寿命、可追溯性、进口依赖和维修能力分析。")
    add_requirement(doc, "FR-MKT-003", "产业链图谱", "系统应呈现上游原材料、核心部件、制造商、经销商、维修服务商、工厂和替代关系。")
    add_requirement(doc, "FR-MKT-004", "市场风险清单", "系统应识别政策、地缘、产能、物流、价格和单一来源风险，并给出影响对象、风险等级和应对建议。")
    add_requirement(doc, "FR-MKT-005", "来源与时效", "外部市场结论应标明来源、发布日期、采集时间和可信等级，过时信息应提示用户。")

    add_heading(doc, "5.4 报价整理与智能比价", 2)
    add_requirement(doc, "FR-RFQ-001", "报价批量解析", "系统应将不同供应商、不同版式的报价自动整理为统一报价明细。")
    add_requirement(doc, "FR-RFQ-002", "可比口径换算", "系统应统一币种、税率、单位、包装、MOQ、运费和到厂条件，计算含税价、未税价及到厂价。")
    add_requirement(doc, "FR-RFQ-003", "缺项与异常检测", "系统应识别报价缺项、规格不一致、异常高低价、有效期不足及与历史价格显著偏离等问题。")
    add_requirement(doc, "FR-RFQ-004", "多维比价", "系统应比较报价、历史采购价、预算价、目标价、市场参考价、交期、质量、服务和风险。")
    add_requirement(doc, "FR-RFQ-005", "招标比价方案", "系统应分别输出最低报价、最低TCO、最高综合得分和最低风险方案，并解释差异。")
    add_requirement(doc, "FR-RFQ-006", "澄清与议价建议", "系统应根据异常项和差异自动生成澄清问题、谈判重点和目标价格建议。")
    add_requirement(doc, "FR-RFQ-007", "比价结果审批", "业务人员可修改口径、权重和结论；系统须记录修改人、修改前后值及原因。")
    add_requirement(doc, "FR-RFQ-008", "疑似围串标识别", "系统可基于报价一致性、文本相似度、联系人及行为模式提示疑似异常，但不得自动认定违规。", "应该")

    add_heading(doc, "5.5 采购价格、TCO与统计模型", 2)
    add_requirement(doc, "FR-COST-001", "成本构成配置", "系统应支持配置采购单价、税费、汇率、运输、保险、仓储、资金、检验、返工、退货、报废、延期和服务损失等成本项。")
    add_requirement(doc, "FR-COST-002", "TCO计算", "系统应按企业确认的口径计算单笔、物料、供应商和项目维度的总拥有成本。")
    add_requirement(doc, "FR-COST-003", "价格趋势分析", "系统应展示历史价格、移动平均、同比、环比、异常波动和价格变化原因。")
    add_requirement(doc, "FR-COST-004", "多元回归模型", "系统应使用价格、延期、服务响应、不良率、返修率、物流及付款成本等多个变量预测实际采购成本。")
    add_requirement(doc, "FR-COST-005", "一元回归教学模式", "系统可提供单变量回归演示，用于培训用户理解变量关系，但不得替代正式多因素成本评价。", "应该")
    add_requirement(doc, "FR-COST-006", "模型诊断", "系统应展示样本量、缺失值处理、异常值、共线性、系数、拟合优度、误差、置信区间和验证结果。")
    add_requirement(doc, "FR-COST-007", "模型版本管理", "模型训练、参数调整、上线、停用和回滚应形成版本记录，生产评分应标记使用的模型版本。")
    add_requirement(doc, "FR-COST-008", "小样本兜底", "历史数据不足时，系统应启用经审批的规则评分或加权评分，并提示其局限性。")
    add_requirement(doc, "FR-COST-009", "招标成本预测", "用户录入或导入投标数据后，系统应预测实际采购成本、给出区间并生成供应商排序。")
    add_requirement(doc, "FR-COST-010", "敏感性分析", "系统应展示权重或关键指标变化对排名和推荐结果的影响。", "应该")

    add_heading(doc, "5.6 需求预测与采购计划", 2)
    add_requirement(doc, "FR-PLAN-001", "预测数据准备", "系统应整合历史领用、销售订单、生产计划、当前库存、在途、退货和替代料数据。")
    add_requirement(doc, "FR-PLAN-002", "基础预测方法", "系统应支持移动平均、指数平滑、季节性分析和简单回归等方法，并可比较误差。")
    add_requirement(doc, "FR-PLAN-003", "预测结果解释", "系统应说明趋势、季节性、异常点、主要驱动因素和置信区间。")
    add_requirement(doc, "FR-PLAN-004", "采购建议量", "系统应结合安全库存、再订货点、MOQ、采购周期、在途、保质期和替代关系生成采购建议。")
    add_requirement(doc, "FR-PLAN-005", "情景模拟", "用户应可模拟需求增长、供应中断、价格上涨、交期延长等情景并比较影响。")
    add_requirement(doc, "FR-PLAN-006", "预警", "系统应生成缺料、积压、呆滞、临期和预测偏差预警。")

    add_heading(doc, "5.7 智能寻源与供应商准入", 2)
    add_requirement(doc, "FR-SRC-001", "寻源任务", "用户应可按物料规格、品类、区域、产能、设备、认证和交期创建寻源条件。")
    add_requirement(doc, "FR-SRC-002", "候选供应商发现", "系统应从企业供应商库及经批准的外部数据源获取候选供应商，并标明数据来源。")
    add_requirement(doc, "FR-SRC-003", "能力匹配", "系统应对候选供应商的产品、工艺、产能、客户、认证、区域和历史绩效进行匹配评分。")
    add_requirement(doc, "FR-SRC-004", "国产替代与第二来源", "系统应提示国产替代、区域替代和第二供应源机会及转换风险。")
    add_requirement(doc, "FR-SRC-005", "初审材料", "系统应生成供应商调查问卷、资料清单、现场审核清单和差距项。")
    add_requirement(doc, "FR-SRC-006", "准入评分", "系统应支持技术、商务、质量、交付、服务、合规和风险多维准入评价。")
    add_requirement(doc, "FR-SRC-007", "长短名单", "系统应形成候选长名单、短名单、淘汰原因和待补充证据。")

    add_heading(doc, "5.8 供应商风险、绩效、关系与价值管理", 2)
    add_requirement(doc, "FR-SUP-001", "供应商360档案", "系统应汇总基本信息、物料、工厂、合同、报价、订单、质量、交付、服务、资质、风险和改善项目。")
    add_requirement(doc, "FR-SUP-002", "风险识别", "系统应识别财务、经营、法律、质量、交付、产能、单一来源、认证、ESG、地缘和物流风险。")
    add_requirement(doc, "FR-SUP-003", "风险预警", "风险事件应形成红黄绿等级、影响物料、影响工厂、建议措施、责任人、截止日期和处理状态。")
    add_requirement(doc, "FR-SUP-004", "绩效指标", "系统应计算准时交付率、来料合格率、成品率、返修率、响应及时度、价格竞争力和改善关闭率。")
    add_requirement(doc, "FR-SUP-005", "绩效周期", "系统应支持月度、季度、年度评价，保存评分依据、人工调整和审批记录。")
    add_requirement(doc, "FR-SUP-006", "供应商分层", "系统应按战略价值与供应风险划分战略型、杠杆型、瓶颈型和常规型供应商。")
    add_requirement(doc, "FR-SUP-007", "关系策略", "系统应基于分层和绩效提出维持、辅导、改善、降级、替换或战略合作建议。")
    add_requirement(doc, "FR-SUP-008", "价值与降本", "系统应跟踪联合研发、材料替代、工艺优化、标准化、合并采购和年度协议等价值项目。")
    add_requirement(doc, "FR-SUP-009", "收益口径", "系统应区分报价降本、成本避免和已实现财务收益，并保存计算依据。")

    add_heading(doc, "5.9 多工厂供货分配与交货路径优化", 2)
    add_requirement(doc, "FR-LOG-001", "供货网络建模", "系统应维护供应商、工厂、仓库、路线、距离、时效、运费、产能、需求和装载限制。")
    add_requirement(doc, "FR-LOG-002", "多目标优化", "系统应支持最低总成本、最短交期、最低风险和最低碳排等目标或组合权重。")
    add_requirement(doc, "FR-LOG-003", "供货分配", "系统应计算供应商向不同工厂的供货数量及约束满足情况。")
    add_requirement(doc, "FR-LOG-004", "路径方案", "系统应比较直送、集货、区域仓和中转仓方案，给出路线、批次和频率建议。")
    add_requirement(doc, "FR-LOG-005", "异常重算", "供应商停产、产能下降、道路中断或工厂需求变化时应支持重新求解。")
    add_requirement(doc, "FR-LOG-006", "算法边界", "路径和分配结果必须由可验证的优化算法产生，大语言模型仅负责解释结果。")

    add_heading(doc, "5.10 AI合同管理", 2)
    add_requirement(doc, "FR-CTR-001", "合同模板库", "系统应支持采购合同、框架协议、质量协议、保密协议等模板的版本、权限和生效管理。")
    add_requirement(doc, "FR-CTR-002", "合同自动创建", "系统应根据供应商、物料、价格、交期、付款和中标结果生成合同初稿。")
    add_requirement(doc, "FR-CTR-003", "条款识别", "系统应识别付款、交付、验收、质保、违约、保密、知识产权、合规和终止条款。")
    add_requirement(doc, "FR-CTR-004", "模板差异审查", "系统应比较供应商合同与企业标准模板，标记缺失、冲突、非标准和高风险条款。")
    add_requirement(doc, "FR-CTR-005", "修改建议", "系统应给出修改建议、风险原因和参考标准条款，并保留法务复核入口。")
    add_requirement(doc, "FR-CTR-006", "业务一致性检查", "合同金额、价格、交期、付款和供应商信息应与报价、中标及审批结果核对。")
    add_requirement(doc, "FR-CTR-007", "履约节点提取", "系统应提取有效期、续签、保证金、价格调整、交付和验收节点并自动提醒。")
    add_requirement(doc, "FR-CTR-008", "合同安全限制", "系统不得将供应商合同、商业秘密或个人信息发送给未批准的大模型或外部服务。")

    add_heading(doc, "5.11 异常、合规与审计", 2)
    add_requirement(doc, "FR-CMP-001", "采购异常检测", "系统应识别拆单、超预算、超授权、无合同、先采购后审批、订单合同价格不一致等异常。")
    add_requirement(doc, "FR-CMP-002", "付款与账户异常", "系统应提示重复订单、重复发票、重复付款及供应商银行账户异常变更。")
    add_requirement(doc, "FR-CMP-003", "关联风险", "系统可提示同一联系人、地址、电话或账户关联多家供应商等可疑关系，并要求人工核验。")
    add_requirement(doc, "FR-CMP-004", "审计日志", "系统应记录登录、查询、上传、下载、数据修改、模型调用、提示词、响应、评分、审批和配置变更。")
    add_requirement(doc, "FR-CMP-005", "审计检索与导出", "授权审计人员可按时间、用户、任务、供应商、合同和操作类型检索并导出日志。")

    add_heading(doc, "5.12 工作台与消息通知", 2)
    add_requirement(doc, "FR-UI-001", "统一工作台", "系统应提供市场洞察、需求计划、寻源比价、成本分析、供应商管理、合同履约六类工作台。")
    add_requirement(doc, "FR-UI-002", "采购驾驶舱", "管理层应可查看采购金额、降本、价格趋势、供应商绩效、风险、合同到期和缺料预警。")
    add_requirement(doc, "FR-UI-003", "任务中心", "用户应可查看待处理、运行中、待复核、已完成、失败和已归档任务。")
    add_requirement(doc, "FR-UI-004", "消息通知", "系统应支持站内消息，并预留邮件、企业微信、钉钉等通知接口。")
    add_requirement(doc, "FR-UI-005", "导出", "经授权用户可导出Excel、Word或PDF结果，导出内容应带生成时间、口径和数据范围。")

    add_heading(doc, "6. 数据需求", 1)
    add_heading(doc, "6.1 核心数据域", 2)
    data_domains = [
        ("DATA-001", "组织与权限", "集团、公司、工厂、采购组织、部门、用户、角色、数据权限"),
        ("DATA-002", "采购主数据", "物料、品类、单位、币种、税率、仓库、工厂、采购员"),
        ("DATA-003", "供应商", "基本信息、联系人、资质、认证、产能、设备、物料能力、风险与分层"),
        ("DATA-004", "交易数据", "需求、询价、报价、招标、合同、订单、收货、入库、发票、付款"),
        ("DATA-005", "质量与服务", "检验、合格、不良、返修、退货、投诉、响应、改善记录"),
        ("DATA-006", "物流", "起点、终点、路线、距离、时效、运费、容量、批次与碳排"),
        ("DATA-007", "知识与文档", "制度、模板、合同、报价、市场资料、行业报告、会议纪要"),
        ("DATA-008", "模型与审计", "数据集、特征、参数、指标、版本、提示词、响应、审批、日志"),
    ]
    add_table(doc, ["编号", "数据域", "主要内容"], data_domains, [1100, 1700, 6560])

    add_heading(doc, "6.2 数据质量与口径要求", 2)
    data_rules = [
        "供应商、物料和组织应具备唯一标识，并建立别名及外部系统编码映射。",
        "价格数据必须明确币种、含税状态、税率、单位、数量阶梯和贸易条件。",
        "质量指标必须统一分母、统计周期、检验范围和是否包含让步接收。",
        "交付指标必须明确承诺日期、变更日期、实际日期和责任归属。",
        "模型训练数据应记录提取时间、清洗规则、排除样本和人工修正。",
        "API导入数据应保留源系统记录标识、源数据更新时间、同步批次、映射版本和处理结果，以支持全链路追溯。",
        "业务删除原则上采用逻辑删除；已进入审批、模型或审计的数据不得直接物理删除。",
    ]
    for item in data_rules:
        add_bullet(doc, item)

    add_heading(doc, "7. 外部接口与集成需求", 1)
    interfaces = [
        ("INT-001", "ERP/SRM", "物料、供应商、需求、询价、报价、合同、订单、收货、付款；支持查询及经批准的结果回写。"),
        ("INT-002", "MES/WMS", "生产计划、领用、库存、在途、批次、仓储和交付数据。"),
        ("INT-003", "QMS", "检验、不良、返修、退货、投诉和供应商改善数据。"),
        ("INT-004", "OA/审批", "任务审批、供应商准入、招标确认、合同审批和消息待办。"),
        ("INT-005", "大模型服务", "支持可替换模型供应商、流式响应、超时、重试、限流、费用统计和敏感数据策略。"),
        ("INT-006", "OCR/文档解析", "支持扫描件、表格、版面和多语言文档识别，并返回页码、位置与置信度。"),
        ("INT-007", "外部市场/风险", "经批准的数据源，包括工商、司法、舆情、价格指数、汇率、物流及行业数据。"),
        ("INT-008", "统一身份认证", "支持企业SSO、OAuth2/OIDC或LDAP等方式，具体以客户环境为准。"),
    ]
    add_table(doc, ["编号", "系统/服务", "接口范围"], interfaces, [1100, 1800, 6460])
    add_requirement(doc, "INT-009", "接口幂等与追踪", "写入类接口应支持幂等键，所有接口调用应记录追踪标识、请求结果和错误信息。")
    add_requirement(doc, "INT-010", "接口降级", "外部大模型、市场数据或业务系统不可用时，系统应提示影响范围，并保留重试或人工处理能力。")
    add_requirement(doc, "INT-011", "数据导入接口规范", "系统API数据导入应优先支持REST/JSON，并预留SOAP/WebService、数据库视图或消息接口的适配能力；接口适配应覆盖分页、超时、限流、重试、签名校验和版本兼容。")
    add_requirement(doc, "INT-012", "数据同步管理", "管理员应可完成接口连通性测试、样例数据预览、字段映射校验、同步计划配置、运行监控、失败重处理和同步任务停启。")

    add_heading(doc, "8. 非功能需求", 1)
    nfr_rows = [
        ("NFR-001", "性能", "常规页面查询P95不高于3秒；普通AI问答应在5秒内返回首段流式内容；长任务进入后台并显示进度。"),
        ("NFR-002", "容量", "一期容量按客户确认的用户数、并发数、文档量、年度采购记录量和模型调用量进行最终核定。"),
        ("NFR-003", "可用性", "核心业务服务目标可用性不低于99.5%；计划维护除外。"),
        ("NFR-004", "安全", "传输使用TLS；敏感数据加密存储；密钥进入专用密钥或环境变量管理，不得写入代码。"),
        ("NFR-005", "模型安全", "支持提示注入防护、工具调用白名单、输出内容检查、敏感信息脱敏及模型供应商隔离策略。"),
        ("NFR-006", "隐私", "个人信息和商业秘密应按最小必要原则处理，支持保留期限、导出控制和删除流程。"),
        ("NFR-007", "审计", "关键操作日志应防篡改并按客户合规要求保留；管理员操作同样必须留痕。"),
        ("NFR-008", "备份恢复", "数据库、对象文件和关键配置应定期备份；RPO、RTO由客户在部署方案中确认。"),
        ("NFR-009", "可维护性", "业务规则、评分权重、提示词、模型、知识库、字典和通知模板应可配置并具备版本。"),
        ("NFR-010", "兼容性", "前端按Vue 2.0实现，支持客户指定的主流桌面浏览器；最低版本在开发文档中确认。"),
        ("NFR-011", "可观测性", "应记录应用、接口、任务、模型调用、费用、错误和性能指标，支持告警和链路追踪。"),
        ("NFR-012", "无障碍与易用性", "关键页面应具备清晰层级、表单校验、空状态、错误提示、加载状态和键盘可操作性。"),
        ("NFR-013", "国际化", "一期至少支持中文；数据层应保留多语言、时区、币种和单位扩展能力。"),
        ("NFR-014", "部署", "支持客户要求的本地化、私有云或指定云环境；容器化及网络隔离方案在开发文档中明确。"),
    ]
    add_table(doc, ["编号", "类别", "要求"], nfr_rows, [1100, 1500, 6760], font_size=9.2)

    add_heading(doc, "9. AI与算法专项要求", 1)
    ai_rules = [
        ("AI-001", "可解释", "每项推荐应展示主要依据、规则、权重、统计指标或引用来源。"),
        ("AI-002", "幻觉控制", "知识问答优先使用授权知识库；证据不足时回答“不足以判断”，不得编造供应商或市场事实。"),
        ("AI-003", "结构化计算", "金额、评分、回归、预测和路径优化由程序、统计库或求解器执行，大模型不得作为唯一计算引擎。"),
        ("AI-004", "模型可替换", "业务功能不应绑定单一大模型厂商；不同模型应通过统一适配层接入。"),
        ("AI-005", "成本控制", "记录模型、令牌数、响应时间、费用、缓存命中和失败原因，支持用户或组织额度。"),
        ("AI-006", "评测", "上线前建立采购问答、文档提取、合同审查和推荐质量评测集，并保留版本化评测结果。"),
        ("AI-007", "人工审批", "AI不得独立完成供应商淘汰、招标定标、合同签署、付款或对外发送正式文件。"),
        ("AI-008", "数据隔离", "不同组织、项目或客户的数据不得在检索、缓存、训练或日志中相互泄露。"),
    ]
    add_table(doc, ["编号", "原则", "要求"], ai_rules, [1100, 1500, 6760])

    add_heading(doc, "10. 关键业务规则", 1)
    add_heading(doc, "10.1 综合采购成本", 2)
    add_para(
        doc,
        "综合采购成本原则上由采购金额、物流成本、资金成本、延期损失、质量损失、服务损失及企业认可的其他成本组成。各成本项的公式、数据来源、缺失值处理和适用范围应由客户确认，并作为可版本化规则维护。",
    )
    add_heading(doc, "10.2 回归模型", 2)
    add_para(
        doc,
        "客户原始需求提出使用一元线性回归评价交期、价格、服务、质量等因素。由于存在多个影响变量，正式业务模型应采用多元线性回归；一元线性回归仅保留为培训演示或单因素探索。模型不得仅凭相关性解释因果关系，生产使用前必须完成样本、稳定性和误差验证。",
    )
    add_para(
        doc,
        "建议目标形式：实际采购成本 = β0 + β1×采购价格 + β2×延期天数 + β3×服务响应时间 + β4×不良率 + β5×返修率 + β6×物流成本 + β7×付款成本。最终变量、单位和损失函数须依据客户数据可得性确认。",
        color=DARK_BLUE,
    )
    add_heading(doc, "10.3 供应商综合评分", 2)
    add_para(
        doc,
        "当数据不足以训练可靠模型时，可使用价格、交付、质量、服务和风险的加权评分。评分权重、归一化方法、缺失值处理、最低准入门槛和一票否决项必须可配置、可审批、可追溯。",
    )

    add_heading(doc, "11. 验收要求", 1)
    acceptance = [
        ("AC-001", "文件解析", "使用客户提供的典型报价和合同样本验证字段提取、来源定位、置信度和人工校正。"),
        ("AC-002", "智能比价", "至少完成币种、税率、单位、运费和关键商务条款的统一比较，并可追溯至原报价。"),
        ("AC-003", "采购问答", "对评测问题返回有来源的答案；无证据问题能明确拒绝或提示补充资料。"),
        ("AC-004", "成本模型", "可完成数据准备、训练、诊断、版本保存、预测和结果解释，并通过客户确认的验证集。"),
        ("AC-005", "需求预测", "可选择基础模型、展示误差、生成采购建议和情景模拟结果。"),
        ("AC-006", "供应商管理", "可查看360档案、绩效、风险、分层、改善和价值项目。"),
        ("AC-007", "合同审查", "可识别关键条款、与模板对比、标记风险、生成建议并保留人工复核记录。"),
        ("AC-008", "权限安全", "跨组织、跨角色和敏感字段权限测试通过，越权访问被拒绝并记录。"),
        ("AC-009", "审计追溯", "关键结论可追踪至原始数据、规则/模型版本、操作人和审批记录。"),
        ("AC-010", "性能稳定", "达到客户最终确认的并发、响应时间、任务成功率和可用性指标。"),
        ("AC-011", "对话式AI助手", "可通过连续对话完成授权数据查找、模型修正申请、任务优化和报告生成；结果有来源，影响性操作有预览、确认、进度、版本和审计记录。"),
        ("AC-012", "系统API数据导入", "至少选取一个客户真实或模拟业务系统接口，完成鉴权、连通性测试、字段映射、全量与增量同步、数据质量校验、异常重试、幂等处理和审计追溯。"),
    ]
    add_table(doc, ["编号", "验收主题", "验收要点"], acceptance, [1100, 1800, 6460])
    add_para(
        doc,
        "详细验收样例、测试数据、量化准确率阈值和性能容量将在客户确认数据条件与一期范围后，写入开发文档及测试方案。",
        color=MUTED,
        size=10,
    )

    add_heading(doc, "12. 边界与不在本期范围", 1)
    out_scope = [
        "完全替代ERP、SRM、MES、WMS、QMS或OA的核心记账和法定审批功能。",
        "未经授权自动对外发送询价、合同、定标通知或供应商处罚决定。",
        "由AI独立作出中标、签约、付款、供应商淘汰等高影响决定。",
        "在一期内建设所有行业的通用外部供应商数据库和实时市场数据平台。",
        "在未获得客户业务数据与口径前承诺回归模型、需求预测或风险模型达到固定准确率。",
        "将企业敏感数据用于公共模型训练，除非客户另行书面授权并完成合规评估。",
    ]
    for item in out_scope:
        add_bullet(doc, item)

    add_heading(doc, "13. 假设、依赖与风险", 1)
    risks = [
        ("数据质量", "历史采购、质量、交期和服务数据可能缺失或口径不一致。", "先完成数据盘点与口径确认，设置规则评分兜底。"),
        ("外部数据", "市场、工商、司法和价格指数可能受授权、费用及更新频率限制。", "需求评审时确定数据源、授权和预算。"),
        ("Vue 2.0", "Vue 2已进入维护期，生态与长期安全维护能力弱于Vue 3。", "按客户要求采用Vue 2.0，同时锁定依赖版本并制定升级预案。"),
        ("模型隐私", "云端大模型可能带来商业秘密与个人信息外发风险。", "配置模型路由、脱敏、私有化选项和数据出境策略。"),
        ("算法可靠性", "小样本、偏差数据或业务变化可能导致模型失真。", "建立评测、监控、人工复核、停用和回滚机制。"),
        ("系统集成", "客户系统接口、鉴权、限流、主数据和网络环境尚未确认。", "在开发前完成接口清单、API文档、样例数据、同步策略和联调环境确认。"),
    ]
    add_table(doc, ["事项", "风险/依赖", "建议措施"], risks, [1500, 3900, 3960], font_size=9.3)

    add_heading(doc, "14. 技术约束与推荐方案（待确认）", 1)
    add_heading(doc, "14.1 前端约束", 2)
    add_para(
        doc,
        "前端采用Vue 2.0。建议配套Vue Router、Vuex、Axios及成熟的Vue 2组件库。由于Vue 2已停止主流维护，开发文档中需明确具体版本锁定、依赖安全扫描、浏览器兼容范围和未来升级路径。",
    )
    add_heading(doc, "14.2 后端推荐", 2)
    add_para(
        doc,
        "推荐采用Python FastAPI作为后端框架。原因包括：适合异步调用大模型和流式输出；内置OpenAPI/Swagger接口文档便于联调；Python生态便于集成OCR、RAG、Pandas、scikit-learn、优化求解器及主流大模型SDK；类型校验清晰，适合构建统一模型适配层和后台任务接口。",
    )
    add_para(
        doc,
        "建议的配套方向为FastAPI + PostgreSQL + Redis + 后台任务队列 + 对象存储 + 向量检索能力。具体采用Celery、RQ或其他任务框架，以及采用pgvector或独立向量数据库，应根据并发、部署方式、数据量和客户基础设施在开发文档中确定。",
        color=DARK_BLUE,
    )
    add_heading(doc, "14.3 选型待确认事项", 2)
    choices = [
        "大模型采用公有云API、私有化模型，还是混合路由。",
        "部署在客户本地、私有云或指定公有云。",
        "是否已有PostgreSQL、MySQL、Redis、对象存储和容器平台等基础设施。",
        "是否需要接入现有ERP、SRM、QMS、OA及统一身份认证。",
        "一期是否需要实时外部市场、工商、司法和价格数据。",
        "一期用户数、并发数、年度采购记录量、文档量和文件大小上限。",
    ]
    for item in choices:
        add_bullet(doc, item)

    add_heading(doc, "15. 待客户确认清单", 1)
    confirmations = [
        ("C-001", "一期范围", "是否按本文“一期”功能作为首个开发版本范围。"),
        ("C-002", "行业优先级", "是否以航空维修业作为首个行业模板，其他制造业作为通用能力。"),
        ("C-003", "回归目标", "“实际采购成本”的精确定义、成本项、统计周期和样本范围。"),
        ("C-004", "质量口径", "成品率、来料合格率、不良率、返修率的业务定义和数据来源。"),
        ("C-005", "供应商评分", "评价维度、权重、门槛、一票否决项及审批角色。"),
        ("C-006", "技术栈", "确认Vue 2.0 + FastAPI方向，以及数据库、缓存、任务和向量检索方案。"),
        ("C-007", "模型策略", "模型供应商、私有化要求、敏感数据规则和预算限制。"),
        ("C-008", "系统集成", "需对接的系统、接口能力、联调环境和数据同步方式。"),
        ("C-009", "部署安全", "部署位置、网络隔离、身份认证、日志留存、备份及灾备要求。"),
        ("C-010", "验收指标", "文档提取准确率、问答评测、预测误差、响应时间和可用性阈值。"),
        ("C-011", "API数据接入", "需接入的系统清单、API文档、鉴权方式、接口地址、限流策略、全量/增量规则、样例数据和联调环境。"),
    ]
    add_table(doc, ["编号", "确认主题", "需确认内容"], confirmations, [1100, 1900, 6360])

    add_heading(doc, "16. 需求变更管理", 1)
    add_para(
        doc,
        "本文件经确认后作为V1.2需求基线。新增模块、改变一期范围、修改业务口径、调整验收阈值或变更技术约束时，应提交需求变更记录，评估对设计、工期、成本、数据、测试和上线计划的影响，经相关责任人确认后纳入新版本。",
    )

    add_heading(doc, "附录A：需求优先级说明", 1)
    priority_rows = [
        ("必须", "Must", "缺失将导致核心业务流程不可用、不能验收或存在重大安全合规风险。"),
        ("应该", "Should", "对效率和体验有明显价值，若一期资源不足可延期，但需保留设计能力。"),
        ("可以", "Could", "增强型或探索型功能，依据价值、数据和预算进入后续版本。"),
    ]
    add_table(doc, ["优先级", "英文", "定义"], priority_rows, [1500, 1500, 6360])

    add_heading(doc, "附录B：培训需求与系统能力映射", 1)
    mapping = [
        ("理解AI在采购中的作用", "AI问答、任务编排、证据引用、人工复核", "FR-AI-*"),
        ("有效提问", "提示词模板、连续追问、任务拆解", "FR-AI-001～003"),
        ("对话式AI助手", "数据查找、上下文控制、模型修正、任务优化、报告版本与执行确认", "FR-AI-008～015"),
        ("数据导入与系统API接入", "文件导入、API同步、映射转换、校验、重试、幂等与审计", "FR-DATA-001、008～013 / INT-001～004、011～012"),
        ("市场环境与产业链分析", "市场专题、航空维修、产业链图谱、风险清单", "FR-MKT-*"),
        ("整理报价与价格分析", "报价解析、口径换算、异常检测、多维比价", "FR-RFQ-*"),
        ("预测需求", "基础预测、误差比较、采购建议、情景模拟", "FR-PLAN-*"),
        ("采购成本与回归模型", "TCO、多元回归、诊断、版本、招标预测", "FR-COST-*"),
        ("供应商管理", "寻源、准入、风险、绩效、关系、价值", "FR-SRC-* / FR-SUP-*"),
        ("多工厂交货路径", "供货网络、多目标优化、分配、路径、重算", "FR-LOG-*"),
        ("合同管理", "模板、生成、审查、一致性、履约提醒", "FR-CTR-*"),
    ]
    add_table(doc, ["培训需求", "系统能力", "需求编号"], mapping, [2600, 5000, 1760], font_size=9.0)

    doc.core_properties.title = "制造业采购AI智能体系统需求规格说明书"
    doc.core_properties.subject = "PEBS Copilot Purchase"
    doc.core_properties.author = "项目需求组"
    doc.core_properties.keywords = "采购AI, 制造业, 系统需求, FastAPI, Vue 2"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_document()
