from io import BytesIO
from uuid import uuid4

from docx import Document
from fastapi.testclient import TestClient

from app.main import app

client = TestClient(app)


def create_contract(headers: dict, suffix: str) -> dict:
    response = client.post(
        "/api/v1/contracts",
        headers=headers,
        json={"title": f"AI审查测试合同-{suffix}", "supplier_id": f"SUP-{suffix}", "supplier_name": "精密零件供应商", "amount": 168000},
    )
    assert response.status_code == 201
    return response.json()


def test_contract_ai_review_version_control_and_submission_gate() -> None:
    buyer = {"X-User-Id": "buyer-review", "X-User-Role": "buyer"}
    suffix = uuid4().hex[:8]
    contract = create_contract(buyer, suffix)
    text = """采购合同
甲方采购轴承一批，精密零件供应商负责供货。合同金额人民币168000元，含税税率13%。
交货期为收到订单后15日，交付地点为甲方工厂。质量要求按图纸执行，甲方在到货后5日内验收。
付款方式为验收合格并收到发票后30日结算。质保期12个月，违约方应承担违约责任。
争议由甲方所在地人民法院诉讼解决。不可抗力发生后3日内书面通知。
"""
    uploaded = client.post(
        f"/api/v1/contracts/{contract['id']}/documents",
        headers=buyer,
        files={"file": ("采购合同V1.txt", text.encode(), "text/plain")},
    )
    assert uploaded.status_code == 201
    document = uploaded.json()
    assert document["version"] == 1
    assert document["reviewed"] is False

    blocked = client.post(f"/api/v1/contracts/{contract['id']}/submit", headers=buyer)
    assert blocked.status_code == 409
    assert "尚未完成AI审查" in blocked.json()["detail"]

    reviewed = client.post(f"/api/v1/contracts/{contract['id']}/ai-review", headers=buyer, json={})
    assert reviewed.status_code == 201
    result = reviewed.json()
    assert result["document_version"] == 1
    assert result["engine"] == "rules"
    assert result["result"]["statistics"]["clause_total"] == 11
    assert result["result"]["clauses"]
    assert result["result"]["disclaimer"]

    workspace = client.get(f"/api/v1/contracts/{contract['id']}/review-workspace")
    assert workspace.status_code == 200
    assert workspace.json()["current_document_reviewed"] is True
    assert workspace.json()["latest_review"]["id"] == result["id"]
    assert client.get(f"/api/v1/contract-documents/{document['id']}/download").status_code == 200
    assert client.delete(f"/api/v1/contract-documents/{document['id']}", headers=buyer).status_code == 409

    history = client.get(f"/api/v1/contracts/{contract['id']}/ai-reviews?page=1&page_size=10").json()
    assert history["total"] == 1
    assert history["items"][0]["document_version"] == 1
    assert client.post(f"/api/v1/contracts/{contract['id']}/submit", headers=buyer).json()["status"] == "pending_approval"


def test_contract_review_docx_upload_new_version_and_permissions() -> None:
    manager = {"X-User-Id": "manager-review", "X-User-Role": "procurement_manager"}
    analyst = {"X-User-Id": "analyst-review", "X-User-Role": "analyst"}
    suffix = uuid4().hex[:8]
    contract = create_contract(manager, suffix)
    invalid = client.post(f"/api/v1/contracts/{contract['id']}/documents", headers=manager, files={"file": ("合同.xls", b"invalid", "application/vnd.ms-excel")})
    assert invalid.status_code == 422

    document = Document()
    document.add_heading("设备零件采购合同", 0)
    document.add_paragraph("精密零件供应商向采购方供应设备零件，合同金额50000元，含税。")
    document.add_paragraph("到货后验收，验收合格后30日付款。违约方承担违约责任，质保期一年。")
    buffer = BytesIO(); document.save(buffer)
    denied = client.post(f"/api/v1/contracts/{contract['id']}/documents", headers=analyst, files={"file": ("合同V1.docx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert denied.status_code == 403
    uploaded = client.post(f"/api/v1/contracts/{contract['id']}/documents", headers=manager, files={"file": ("合同V1.docx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert uploaded.status_code == 201
    assert uploaded.json()["extraction_mode"] == "docx"
    assert client.post(f"/api/v1/contracts/{contract['id']}/ai-review", headers=analyst, json={}).status_code == 403

    second = client.post(f"/api/v1/contracts/{contract['id']}/documents", headers=manager, files={"file": ("合同V2.txt", ("采购合同V2。精密零件供应商供货，合同金额50000元。" * 3).encode(), "text/plain")})
    assert second.status_code == 201
    assert second.json()["version"] == 2
    assert client.delete(f"/api/v1/contract-documents/{second.json()['id']}", headers=manager).status_code == 204
    workspace = client.get(f"/api/v1/contracts/{contract['id']}/review-workspace").json()
    assert workspace["documents"][0]["version"] == 1


def test_contract_pdf_upload_extracts_chinese_text(monkeypatch) -> None:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas

    buyer = {"X-User-Id": "buyer-pdf", "X-User-Role": "buyer"}
    contract = create_contract(buyer, uuid4().hex[:8])
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    buffer = BytesIO()
    document = canvas.Canvas(buffer)
    document.setFont("STSong-Light", 12)
    document.drawString(72, 760, "采购合同 合同金额100000元 交付期为30日 付款验收后结算")
    document.save()
    response = client.post(
        f"/api/v1/contracts/{contract['id']}/documents",
        headers=buyer,
        files={"file": ("采购合同.pdf", buffer.getvalue(), "application/pdf")},
    )
    assert response.status_code == 201
    assert response.json()["extraction_mode"] == "pdf_text"
    assert response.json()["text_length"] >= 30
