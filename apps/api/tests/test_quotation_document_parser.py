import io
import zipfile

from docx import Document

from app.domain.rfqs import RFQLine
from app.services.quotation_document_parser import parse_quotation


LINES = [RFQLine(material_code="MAT-01", material_name="测试件", quantity=10, unit="件")]


def test_docx_table_extraction():
    document = Document()
    table = document.add_table(rows=2, cols=4)
    for cell, value in zip(table.rows[0].cells, ["物料编码", "未税单价", "税率", "物流费"], strict=True):
        cell.text = value
    for cell, value in zip(table.rows[1].cells, ["MAT-01", "18.6", "13%", "5"], strict=True):
        cell.text = value
    document.add_paragraph("交期：7天；有效期：60天")
    stream = io.BytesIO()
    document.save(stream)
    result = parse_quotation(stream.getvalue(), "报价.docx", LINES)
    assert result["confidence"] == 1
    assert result["lines"][0]["unit_price"] == 18.6
    assert result["delivery_days"] == 7
    assert result["validity_days"] == 60


def test_minimal_xlsx_extraction():
    shared = """<?xml version='1.0' encoding='UTF-8'?><sst xmlns='http://schemas.openxmlformats.org/spreadsheetml/2006/main'><si><t>物料编码</t></si><si><t>未税单价</t></si><si><t>税率</t></si><si><t>物流费</t></si><si><t>MAT-01</t></si></sst>"""
    sheet = """<?xml version='1.0' encoding='UTF-8'?><worksheet xmlns='http://schemas.openxmlformats.org/spreadsheetml/2006/main'><sheetData><row r='1'><c t='s'><v>0</v></c><c t='s'><v>1</v></c><c t='s'><v>2</v></c><c t='s'><v>3</v></c></row><row r='2'><c t='s'><v>4</v></c><c><v>22.5</v></c><c><v>0.13</v></c><c><v>8</v></c></row></sheetData></worksheet>"""
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("xl/sharedStrings.xml", shared)
        archive.writestr("xl/worksheets/sheet1.xml", sheet)
    result = parse_quotation(stream.getvalue(), "报价.xlsx", LINES)
    assert result["confidence"] == 1
    assert result["lines"][0]["unit_price"] == 22.5
    assert result["lines"][0]["logistics_cost"] == 8

