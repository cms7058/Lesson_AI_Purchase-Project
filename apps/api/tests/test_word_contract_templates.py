from io import BytesIO
from pathlib import Path
from uuid import uuid4

from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import select

from app.core.database import SessionLocal
from app.domain.persistence import (
    BusinessTemplateRecord,
    ContractElementRecord,
    ContractFieldDefinitionRecord,
    ContractRecord,
    DocumentTemplateLinkRecord,
    GeneratedBusinessDocumentRecord,
    ManagedWordTemplateRecord,
    WordTemplateFileRecord,
)
from app.main import app

client = TestClient(app)
HEADERS = {"X-User-Role": "procurement_manager", "X-User-Id": "word-template-tester"}


def _template_bytes() -> bytes:
    document = Document()
    document.add_heading("采购合同 {{ contract_no }}", 0)
    document.add_paragraph("供应商：{{ supplier_name }}")
    document.add_paragraph("项目：{{ project_name }}")
    table = document.add_table(rows=1, cols=5)
    for cell, value in zip(table.rows[0].cells, ["序号", "编码", "名称", "数量", "含税金额"], strict=True):
        cell.text = value
    start = table.add_row()
    start.cells[0].text = "{%tr for item in items %}"
    row = table.add_row()
    for cell, value in zip(
        row.cells,
        ["{{ item.index }}", "{{ item.material_code }}", "{{ item.material_name }}", "{{ item.quantity }}", "{{ item.line_total }}"],
        strict=True,
    ):
        cell.text = value
    end = table.add_row()
    end.cells[0].text = "{%tr endfor %}"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_downloadable_sample_is_a_valid_docx():
    response = client.get("/api/v1/templates/word-contract/sample", headers=HEADERS)
    assert response.status_code == 200
    document = Document(BytesIO(response.content))
    assert "contract_no" in "\n".join(paragraph.text for paragraph in document.paragraphs)


def test_upload_validate_fill_and_export_word_contract():
    tag = uuid4().hex[:8]
    upload = client.post(
        "/api/v1/templates/word-contract",
        headers=HEADERS,
        data={"name": f"Word合同模板{tag}", "version": "1.0"},
        files={"file": ("采购合同模板.docx", _template_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert upload.status_code == 201, upload.text
    template = upload.json()
    assert template["engine"] == "docx_v1"
    assert template["validation_status"] == "validated"

    validation = client.get(f"/api/v1/templates/{template['id']}/validation", headers=HEADERS)
    assert validation.status_code == 200
    assert "contract_no" in validation.json()["placeholders"]
    assert "items" in validation.json()["placeholders"]

    contract = client.post(
        "/api/v1/contracts",
        headers=HEADERS,
        json={
            "title": f"Word生成测试合同{tag}",
            "supplier_id": "QA-SUP",
            "supplier_name": "Word测试供应商",
            "amount": 1,
            "template_id": template["id"],
            "elements": {"project_name": "智能工厂扩建项目"},
            "items": [
                {"material_code": "M-001", "material_name": "伺服电机", "quantity": 2, "unit": "台", "unit_price": 100, "tax_rate": 13},
                {"material_code": "M-002", "material_name": "联轴器", "quantity": 3, "unit": "件", "unit_price": 50, "tax_rate": 13},
            ],
        },
    )
    assert contract.status_code == 201, contract.text
    contract_data = contract.json()
    assert float(contract_data["amount"]) == 395.5
    assert len(contract_data["items"]) == 2

    exported = client.post(
        f"/api/v1/documents/contract/{contract_data['id']}/export",
        headers=HEADERS,
        json={"template_id": template["id"], "output_format": "docx"},
    )
    assert exported.status_code == 200, exported.text
    path = Path("generated_documents") / exported.json()["file_name"]
    assert path.is_file()
    generated = Document(path)
    body = "\n".join(paragraph.text for paragraph in generated.paragraphs)
    body += "\n" + "\n".join(cell.text for table in generated.tables for row in table.rows for cell in row.cells)
    assert contract_data["contract_no"] in body
    assert "智能工厂扩建项目" in body
    assert "M-001" in body and "M-002" in body

    with SessionLocal.begin() as db:
        generated_records = list(db.scalars(select(GeneratedBusinessDocumentRecord).where(GeneratedBusinessDocumentRecord.source_id == contract_data["id"])))
        for record in generated_records:
            Path(record.storage_path).unlink(missing_ok=True)
            db.delete(record)
        for link in db.scalars(select(DocumentTemplateLinkRecord).where(DocumentTemplateLinkRecord.document_id == contract_data["id"])):
            db.delete(link)
        elements = db.get(ContractElementRecord, contract_data["id"])
        if elements:
            db.delete(elements)
        db.delete(db.get(ContractRecord, contract_data["id"]))
        word_file = db.scalar(select(WordTemplateFileRecord).where(WordTemplateFileRecord.template_id == template["id"]))
        Path(word_file.storage_path).unlink(missing_ok=True)
        db.delete(word_file)
        db.delete(db.get(BusinessTemplateRecord, template["id"]))


def test_custom_field_drives_managed_template_and_contract_rendering():
    tag = uuid4().hex[:8]
    field = client.post(
        "/api/v1/contract-fields", headers=HEADERS,
        json={"code": f"custom_clause_{tag}", "name": "定制服务条款", "category": "服务", "data_type": "textarea", "required": True},
    )
    assert field.status_code == 201, field.text
    field_data = field.json()
    managed = client.post(
        "/api/v1/contract-fields/generate-template", headers=HEADERS,
        json={"name": f"动态合同模板{tag}", "version": "1.0", "field_ids": [field_data["id"]]},
    )
    assert managed.status_code == 201, managed.text
    template = managed.json()
    assert f"custom_clause_{tag}" in template["placeholders"]

    contract = client.post(
        "/api/v1/contracts", headers=HEADERS,
        json={"title": f"动态字段合同{tag}", "supplier_id": "QA-SUP", "supplier_name": "动态字段供应商", "amount": 100, "template_id": template["id"], "elements": {f"custom_clause_{tag}": "四小时内到场服务"}},
    )
    assert contract.status_code == 201, contract.text
    contract_data = contract.json()
    preflight = client.get(
        f"/api/v1/documents/contract/{contract_data['id']}/preflight",
        headers=HEADERS,
        params={"template_id": template["id"]},
    )
    assert preflight.status_code == 200, preflight.text
    assert preflight.json()["valid"] is True
    export = client.post(
        f"/api/v1/documents/contract/{contract_data['id']}/export", headers=HEADERS,
        json={"template_id": template["id"], "output_format": "docx"},
    )
    assert export.status_code == 200, export.text
    output_path = Path("generated_documents") / export.json()["file_name"]
    assert "四小时内到场服务" in "\n".join(item.text for item in Document(output_path).paragraphs)
    history = client.get(f"/api/v1/documents/contract/{contract_data['id']}/history", headers=HEADERS)
    assert history.status_code == 200
    assert history.json()[0]["template_version"] == "1.0"
    assert len(history.json()[0]["sha256"]) == 64

    changed = client.patch(f"/api/v1/contract-fields/{field_data['id']}", headers=HEADERS, json={"name": "更新后的服务条款"})
    assert changed.status_code == 200
    original = client.get(f"/api/v1/templates/{template['id']}/file", headers=HEADERS)
    assert "更新后的服务条款" in "\n".join(item.text for item in Document(BytesIO(original.content)).paragraphs)

    with SessionLocal.begin() as db:
        for generated in db.scalars(select(GeneratedBusinessDocumentRecord).where(GeneratedBusinessDocumentRecord.source_id == contract_data["id"])):
            Path(generated.storage_path).unlink(missing_ok=True); db.delete(generated)
        for link in db.scalars(select(DocumentTemplateLinkRecord).where(DocumentTemplateLinkRecord.document_id == contract_data["id"])):
            db.delete(link)
        elements = db.get(ContractElementRecord, contract_data["id"])
        if elements: db.delete(elements)
        db.delete(db.get(ContractRecord, contract_data["id"]))
        db.delete(db.get(ManagedWordTemplateRecord, template["id"]))
        word_file = db.scalar(select(WordTemplateFileRecord).where(WordTemplateFileRecord.template_id == template["id"]))
        Path(word_file.storage_path).unlink(missing_ok=True); db.delete(word_file)
        db.delete(db.get(BusinessTemplateRecord, template["id"]))
        db.delete(db.get(ContractFieldDefinitionRecord, field_data["id"]))
