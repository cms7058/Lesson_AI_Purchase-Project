import hashlib
import json
import re
import zipfile
from io import BytesIO
from pathlib import Path
from uuid import uuid4

from docx import Document
from docxtpl import DocxTemplate
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.domain.persistence import (
    BusinessTemplateRecord,
    ContractFieldDefinitionRecord,
    ManagedWordTemplateRecord,
    WordTemplateFileRecord,
)

TEMPLATE_DIR = Path("private_word_templates")
MAX_FILE_SIZE = 15 * 1024 * 1024
MAX_UNCOMPRESSED_SIZE = 50 * 1024 * 1024
MAX_ZIP_ENTRIES = 2000

KNOWN_ROOTS = {
    "contract_no", "document_no", "contract_title", "title", "contract_type", "currency",
    "amount", "total_amount", "total_amount_uppercase", "effective_date", "expiry_date",
    "signing_date", "project_name", "project_code", "rfq_no", "order_no", "requisition_no",
    "supplier_id", "supplier_name", "buyer", "supplier", "company", "items", "payment_milestones",
    "delivery_milestones", "delivery_address", "transport_method", "acceptance_standard",
    "quality_standard", "warranty_period", "payment_terms", "invoice_type", "confidentiality",
    "intellectual_property", "breach_liability", "force_majeure", "termination_terms",
    "dispute_resolution", "jurisdiction", "company_logo", "buyer_seal", "supplier_seal",
    "created_date", "created_by", "remarks",
}


class WordTemplateError(ValueError):
    pass


def _extract_expressions(data: bytes) -> tuple[list[str], set[str]]:
    expressions: set[str] = set()
    loop_variables: set[str] = set()
    with zipfile.ZipFile(BytesIO(data)) as archive:
        xml = "\n".join(
            archive.read(name).decode("utf-8", errors="ignore")
            for name in archive.namelist()
            if name.startswith("word/") and name.endswith(".xml")
        )
    # Remove Word XML tags because Word may split one placeholder across runs.
    text = re.sub(r"<[^>]+>", "", xml)
    for match in re.findall(r"\{\{\s*(.+?)\s*\}\}", text):
        expressions.add(match.strip())
    for match in re.findall(r"\{%[a-z]*\s+for\s+([a-zA-Z_]\w*)\s+in\s+([a-zA-Z_]\w*)", text):
        loop_variables.add(match[0])
        expressions.add(match[1])
    return sorted(expressions), loop_variables


def validate_docx_bytes(data: bytes, filename: str, custom_roots: set[str] | None = None) -> dict:
    if not filename.lower().endswith(".docx"):
        raise WordTemplateError("仅支持未加密的 .docx 文件")
    if not data:
        raise WordTemplateError("上传文件为空")
    if len(data) > MAX_FILE_SIZE:
        raise WordTemplateError("Word模板不能超过15MB")
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            names = archive.namelist()
            if "word/document.xml" not in names or "[Content_Types].xml" not in names:
                raise WordTemplateError("文件不是有效的Word DOCX模板")
            if len(names) > MAX_ZIP_ENTRIES:
                raise WordTemplateError("模板内部文件数量异常")
            if sum(item.file_size for item in archive.infolist()) > MAX_UNCOMPRESSED_SIZE:
                raise WordTemplateError("模板解压后体积超过50MB")
            lowered = {name.lower() for name in names}
            if "word/vbaproject.bin" in lowered or any("/embeddings/" in name for name in lowered):
                raise WordTemplateError("模板不能包含宏或嵌入对象")
            relationships = "".join(
                archive.read(name).decode("utf-8", errors="ignore")
                for name in names if name.endswith(".rels")
            )
            if 'TargetMode="External"' in relationships:
                raise WordTemplateError("模板不能包含外部文件或网络链接")
        template = DocxTemplate(BytesIO(data))
        template.init_docx()
        expressions, loop_variables = _extract_expressions(data)
    except WordTemplateError:
        raise
    except (zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
        raise WordTemplateError("DOCX结构损坏或文件已加密") from exc
    known_roots = KNOWN_ROOTS | (custom_roots or set())
    unknown = []
    for expression in expressions:
        root = re.split(r"[.\[| (]", expression, maxsplit=1)[0]
        if root and root not in known_roots and root not in loop_variables:
            unknown.append(expression)
    warnings = []
    if not expressions:
        warnings.append("模板中未识别到可填充字段")
    if unknown:
        warnings.append("部分字段不在标准合同字段字典中，生成前需要补充映射")
    return {
        "valid": not unknown,
        "status": "validated" if not unknown else "needs_mapping",
        "placeholders": expressions,
        "unknown_placeholders": sorted(unknown),
        "warnings": warnings,
    }


def create_word_template(
    db: Session, *, name: str, version: str, data: bytes, filename: str, created_by: str
) -> BusinessTemplateRecord:
    custom_roots = set(db.scalars(select(ContractFieldDefinitionRecord.code).where(ContractFieldDefinitionRecord.active.is_(True))))
    report = validate_docx_bytes(data, filename, custom_roots)
    TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
    template = BusinessTemplateRecord(
        name=name.strip(), template_type="contract", version=version.strip(),
        content="Word合同模板（字段由上传的DOCX文件定义）", status="active", created_by=created_by,
    )
    db.add(template)
    db.flush()
    path = TEMPLATE_DIR / f"{template.id}-{uuid4().hex[:8]}.docx"
    path.write_bytes(data)
    db.add(WordTemplateFileRecord(
        template_id=template.id,
        original_filename=Path(filename).name[:255],
        storage_path=str(path),
        file_size=len(data),
        sha256=hashlib.sha256(data).hexdigest(),
        placeholders_json=json.dumps(report["placeholders"], ensure_ascii=False),
        validation_json=json.dumps(report, ensure_ascii=False),
        validation_status=report["status"],
    ))
    db.flush()
    db.refresh(template)
    return template


def get_word_file(db: Session, template_id: str) -> WordTemplateFileRecord | None:
    return db.scalar(select(WordTemplateFileRecord).where(WordTemplateFileRecord.template_id == template_id))


def validation_report(record: WordTemplateFileRecord) -> dict:
    try:
        return json.loads(record.validation_json)
    except json.JSONDecodeError:
        return {"valid": False, "status": "invalid", "placeholders": [], "unknown_placeholders": [], "warnings": ["校验报告损坏"]}


def revalidate_all_word_templates(db: Session) -> None:
    custom_roots = set(db.scalars(select(ContractFieldDefinitionRecord.code).where(ContractFieldDefinitionRecord.active.is_(True))))
    for record in db.scalars(select(WordTemplateFileRecord)):
        path = Path(record.storage_path)
        if not path.is_file():
            report = {"valid": False, "status": "missing_file", "placeholders": [], "unknown_placeholders": [], "warnings": ["模板源文件已丢失"]}
        else:
            try:
                report = validate_docx_bytes(path.read_bytes(), record.original_filename, custom_roots)
            except WordTemplateError as error:
                report = {"valid": False, "status": "invalid", "placeholders": [], "unknown_placeholders": [], "warnings": [str(error)]}
        record.placeholders_json = json.dumps(report["placeholders"], ensure_ascii=False)
        record.validation_json = json.dumps(report, ensure_ascii=False)
        record.validation_status = report["status"]


def build_managed_contract_template(fields: list[ContractFieldDefinitionRecord]) -> bytes:
    document = Document()
    document.add_heading("采购合同 {{ contract_no }}", level=0)
    document.add_paragraph("甲方：{{ buyer.company_name }}")
    document.add_paragraph("乙方：{{ supplier.company_name }}")
    for field in fields:
        document.add_paragraph(f"{field.name}：{{{{ {field.code} }}}}")
    table = document.add_table(rows=1, cols=7)
    headers = ["序号", "物料编码", "物料名称", "规格", "数量", "未税单价", "含税金额"]
    for cell, value in zip(table.rows[0].cells, headers, strict=True):
        cell.text = value
    table.add_row().cells[0].text = "{%tr for item in items %}"
    values = ["{{ item.index }}", "{{ item.material_code }}", "{{ item.material_name }}", "{{ item.specification }}", "{{ item.quantity }} {{ item.unit }}", "{{ item.unit_price }}", "{{ item.line_total }}"]
    for cell, value in zip(table.add_row().cells, values, strict=True):
        cell.text = value
    table.add_row().cells[0].text = "{%tr endfor %}"
    document.add_paragraph("合同含税总额：{{ currency }} {{ total_amount }}")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def sync_managed_word_templates(db: Session) -> None:
    for managed in db.scalars(select(ManagedWordTemplateRecord)):
        word_file = get_word_file(db, managed.template_id)
        if word_file is None:
            continue
        field_ids = json.loads(managed.field_ids_json or "[]")
        query = select(ContractFieldDefinitionRecord).where(ContractFieldDefinitionRecord.active.is_(True))
        if field_ids:
            query = query.where(ContractFieldDefinitionRecord.id.in_(field_ids))
        fields = list(db.scalars(query.order_by(ContractFieldDefinitionRecord.sort_order)))
        data = build_managed_contract_template(fields)
        Path(word_file.storage_path).write_bytes(data)
        word_file.file_size = len(data)
        word_file.sha256 = hashlib.sha256(data).hexdigest()
