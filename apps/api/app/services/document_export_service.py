import hashlib
import json
import shutil
import subprocess
import tempfile
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from uuid import uuid4

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docxtpl import DocxTemplate, InlineImage
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.domain.persistence import (
    BusinessTemplateRecord,
    CompanyProfileRecord,
    ContractElementRecord,
    ContractFieldDefinitionRecord,
    ContractRecord,
    GeneratedBusinessDocumentRecord,
    PurchaseOrderRecord,
    QuotationRecord,
    SupplierRecord,
    WordTemplateFileRecord,
)

OUTPUT_DIR = Path("generated_documents")


def _json_default(value):
    if isinstance(value, (Decimal, datetime)):
        return str(value)
    return str(value)


def _contract_context(db: Session, template: DocxTemplate, contract: ContractRecord, profile: CompanyProfileRecord) -> dict:
    extension = db.get(ContractElementRecord, contract.id)
    stored = json.loads(extension.data_json) if extension and extension.data_json else {}
    elements = stored.get("elements", {})
    items = []
    for index, source in enumerate(stored.get("items", []), start=1):
        item = dict(source)
        quantity = Decimal(str(item.get("quantity") or 0))
        unit_price = Decimal(str(item.get("unit_price") or 0))
        tax_rate = Decimal(str(item.get("tax_rate") or 0))
        if tax_rate > 1:
            tax_rate /= Decimal(100)
        item.update({
            "index": index,
            "tax_rate": f"{tax_rate * 100:g}%",
            "tax_amount": f"{quantity * unit_price * tax_rate:.2f}",
            "line_total": f"{quantity * unit_price * (Decimal(1) + tax_rate):.2f}",
        })
        items.append(item)
    supplier_record = db.get(SupplierRecord, contract.supplier_id)
    if supplier_record is None:
        supplier_record = db.scalar(select(SupplierRecord).where(SupplierRecord.code == contract.supplier_id))
    buyer = {
        "company_name": profile.company_name or "",
        "address": profile.address or "",
        "contact": profile.contact or "",
    }
    supplier = {
        "supplier_id": contract.supplier_id,
        "company_name": contract.supplier_name,
        "unified_credit_code": supplier_record.unified_credit_code if supplier_record else "",
        "address": supplier_record.address if supplier_record else "",
        "contact": supplier_record.contact if supplier_record else "",
        "phone": supplier_record.phone if supplier_record else "",
        "email": supplier_record.email if supplier_record else "",
    }
    context = {
        **elements,
        "document_no": contract.contract_no,
        "contract_no": contract.contract_no,
        "contract_title": contract.title,
        "title": contract.title,
        "supplier_id": contract.supplier_id,
        "supplier_name": contract.supplier_name,
        "amount": f"{contract.amount:.2f}",
        "total_amount": f"{contract.amount:.2f}",
        "currency": contract.currency,
        "effective_date": str(contract.effective_date or ""),
        "expiry_date": str(contract.expiry_date or ""),
        "created_date": str(contract.created_at.date()),
        "created_by": contract.created_by,
        "buyer": buyer,
        "supplier": supplier,
        "company": buyer,
        "items": items,
        "contract": {
            "contract_no": contract.contract_no,
            "title": contract.title,
            "amount": f"{contract.amount:.2f}",
            "currency": contract.currency,
            "effective_date": str(contract.effective_date or ""),
            "expiry_date": str(contract.expiry_date or ""),
        },
    }
    for field in db.scalars(select(ContractFieldDefinitionRecord).where(ContractFieldDefinitionRecord.active.is_(True))):
        if field.code in elements:
            context[field.code] = elements[field.code]
            continue
        value = None
        if field.source_path and field.source_path != "manual":
            value = context
            for segment in field.source_path.split("."):
                value = value.get(segment) if isinstance(value, dict) else None
                if value is None:
                    break
        context[field.code] = value if value not in (None, "") else (field.default_value or "")
    logo_path = Path((profile.logo_path or "").replace("/files/assets/", "generated_assets/"))
    if logo_path.is_file():
        context["company_logo"] = InlineImage(template, str(logo_path), width=Inches(0.7))
    else:
        context["company_logo"] = ""
    return context


def _values(source_type: str, record: object) -> dict[str, str]:
    if source_type == "order":
        order = record
        line = order.lines[0] if order.lines else None
        return {"document_no": order.order_no, "order_no": order.order_no, "supplier_name": order.supplier_name, "supplier_id": order.supplier_id, "factory_code": order.factory_code, "currency": order.currency, "material_code": line.material_code if line else "", "material_name": line.material_name if line else "", "quantity": str(line.quantity) if line else "", "unit": line.unit if line else "", "unit_price": str(line.unit_price) if line else ""}
    if source_type == "quotation":
        quote = record
        line = quote.lines[0] if quote.lines else None
        return {"document_no": quote.quotation_no, "quotation_no": quote.quotation_no, "supplier_name": quote.supplier_name, "supplier_id": quote.supplier_id, "currency": quote.currency, "delivery_days": str(quote.delivery_days), "material_code": line.material_code if line else "", "material_name": line.material_name if line else "", "quantity": str(line.quantity) if line else "", "unit": line.unit if line else "", "unit_price": str(line.unit_price) if line else ""}
    contract = record
    return {"document_no": contract.contract_no, "contract_no": contract.contract_no, "contract_title": contract.title, "supplier_name": contract.supplier_name, "supplier_id": contract.supplier_id, "currency": contract.currency, "amount": str(contract.amount), "effective_date": str(contract.effective_date or ""), "expiry_date": str(contract.expiry_date or "")}


def _render(content: str, values: dict[str, str]) -> str:
    for key, value in values.items():
        content = content.replace(f"{{{{{key}}}}}", value)
    return content


class DocumentExportService:
    def preflight(self, db: Session, source_type: str, source_id: str, template_id: str) -> dict:
        template = db.get(BusinessTemplateRecord, template_id)
        record_type = {"order": PurchaseOrderRecord, "quotation": QuotationRecord, "contract": ContractRecord}[source_type]
        record = db.get(record_type, source_id)
        if template is None or record is None or template.template_type != source_type:
            raise ValueError("未找到模板或业务单据")
        result = {
            "valid": template.status == "active",
            "template_id": template.id,
            "template_name": template.name,
            "template_version": template.version,
            "engine": "text_v1",
            "validation_status": "validated",
            "placeholders": [],
            "unresolved": [],
            "missing_required": [],
            "warnings": [],
            "pdf_available": bool(shutil.which("libreoffice") or shutil.which("soffice")),
        }
        word_file = db.scalar(select(WordTemplateFileRecord).where(WordTemplateFileRecord.template_id == template_id))
        if not word_file:
            if template.status != "active":
                result["warnings"].append("模板未启用")
            return result
        result["engine"] = "docx_v1"
        result["validation_status"] = word_file.validation_status
        result["placeholders"] = json.loads(word_file.placeholders_json or "[]")
        if word_file.validation_status != "validated" or not Path(word_file.storage_path).is_file():
            result["valid"] = False
            result["warnings"].append("Word模板未通过校验或源文件已丢失")
            return result
        source = DocxTemplate(word_file.storage_path)
        profile = db.get(CompanyProfileRecord, 1) or CompanyProfileRecord(id=1)
        context = _contract_context(db, source, record, profile)
        result["unresolved"] = sorted(source.get_undeclared_template_variables(context=context))
        placeholder_set = set(result["placeholders"])
        result["missing_required"] = [
            field.name
            for field in db.scalars(
                select(ContractFieldDefinitionRecord).where(
                    ContractFieldDefinitionRecord.active.is_(True),
                    ContractFieldDefinitionRecord.required.is_(True),
                )
            )
            if field.code in placeholder_set and not context.get(field.code)
        ]
        if "items" in placeholder_set and not context["items"]:
            result["warnings"].append("模板包含采购明细表，但当前合同尚未添加物料")
        if not result["pdf_available"]:
            result["warnings"].append("当前运行环境未安装LibreOffice，只能生成Word")
        result["valid"] = result["valid"] and not result["unresolved"] and not result["missing_required"]
        return result

    def export(self, db: Session, source_type: str, source_id: str, template_id: str, output_format: str, created_by: str = "system") -> tuple[str, str]:
        template = db.get(BusinessTemplateRecord, template_id)
        record_type = {"order": PurchaseOrderRecord, "quotation": QuotationRecord, "contract": ContractRecord}[source_type]
        record = db.get(record_type, source_id)
        if template is None or record is None or template.template_type != source_type:
            raise ValueError("未找到模板或业务单据")
        if source_type in {"order", "quotation"}:
            db.refresh(record, attribute_names=["lines"])
        profile = db.get(CompanyProfileRecord, 1) or CompanyProfileRecord(id=1)
        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        stem = f"{source_type}-{source_id[:8]}-{uuid4().hex[:8]}"
        word_file = db.scalar(select(WordTemplateFileRecord).where(WordTemplateFileRecord.template_id == template_id))
        snapshot: dict = _values(source_type, record)
        if word_file:
            if source_type != "contract":
                raise ValueError("当前Word模板仅支持合同生成")
            if word_file.validation_status != "validated" or not Path(word_file.storage_path).is_file():
                raise ValueError("Word模板未通过校验或源文件已丢失")
            docx_path = OUTPUT_DIR / f"{stem}.docx"
            source = DocxTemplate(word_file.storage_path)
            context = _contract_context(db, source, record, profile)
            placeholders = set(json.loads(word_file.placeholders_json or "[]"))
            required_missing = [
                field.name
                for field in db.scalars(
                    select(ContractFieldDefinitionRecord).where(
                        ContractFieldDefinitionRecord.active.is_(True),
                        ContractFieldDefinitionRecord.required.is_(True),
                    )
                )
                if field.code in placeholders and not context.get(field.code)
            ]
            if required_missing:
                raise ValueError(f"必填合同要素缺失：{', '.join(required_missing)}")
            missing = sorted(source.get_undeclared_template_variables(context=context))
            if missing:
                raise ValueError(f"合同要素缺失：{', '.join(missing)}")
            try:
                source.render(context)
                source.save(docx_path)
            except Exception as error:
                raise ValueError(f"Word模板渲染失败：{error}") from error
            snapshot = context
            if output_format == "docx":
                path = docx_path
            else:
                path = self._convert_to_pdf(docx_path)
        else:
            content = _render(template.content, snapshot)
            if output_format == "docx":
                path = OUTPUT_DIR / f"{stem}.docx"
                self._docx(path, profile, source_type, content)
            else:
                path = OUTPUT_DIR / f"{stem}.pdf"
                self._pdf(path, profile, source_type, content)
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        db.add(GeneratedBusinessDocumentRecord(
            source_type=source_type,
            source_id=source_id,
            template_id=template.id,
            template_version=template.version,
            output_format=output_format,
            file_name=path.name,
            storage_path=str(path),
            sha256=digest,
            source_snapshot_json=json.dumps(snapshot, ensure_ascii=False, default=_json_default),
            created_by=created_by,
        ))
        db.flush()
        return path.name, f"/files/documents/{path.name}"

    def _convert_to_pdf(self, docx_path: Path) -> Path:
        executable = shutil.which("libreoffice") or shutil.which("soffice")
        if executable is None:
            raise ValueError("当前环境未安装LibreOffice，Word已生成但无法转换PDF")
        with tempfile.TemporaryDirectory(prefix="contract-pdf-") as temp_dir:
            result = subprocess.run(
                [executable, "--headless", "--convert-to", "pdf", "--outdir", temp_dir, str(docx_path.resolve())],
                capture_output=True, text=True, timeout=120, check=False,
            )
            generated = Path(temp_dir) / f"{docx_path.stem}.pdf"
            if result.returncode != 0 or not generated.is_file():
                detail = (result.stderr or result.stdout or "转换程序未生成文件").strip()
                raise ValueError(f"PDF转换失败：{detail[:300]}")
            target = docx_path.with_suffix(".pdf")
            shutil.copyfile(generated, target)
            return target

    def _docx(self, path: Path, profile: CompanyProfileRecord, source_type: str, content: str) -> None:
        company_name = profile.company_name or "示例制造集团"
        short_name = profile.short_name or "AI助力"
        doc = Document()
        normal = doc.styles["Normal"]
        normal.font.name = "STSong"
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), "STSong")
        section = doc.sections[0]
        section.top_margin = section.bottom_margin = Inches(0.8)
        section.left_margin = section.right_margin = Inches(0.85)
        header = section.header.paragraphs[0]
        if profile.logo_path:
            local_logo = Path(profile.logo_path.replace("/files/assets/", "generated_assets/"))
            if local_logo.exists():
                header.add_run().add_picture(str(local_logo), width=Inches(0.45))
        head = header.add_run(f"  {company_name}")
        head.bold = True; head.font.size = Pt(11); head.font.name = "STSong"; head._element.rPr.rFonts.set(qn("w:eastAsia"), "STSong")
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run({"order": "采购订单", "quotation": "询价单", "contract": "采购合同"}[source_type])
        run.bold = True; run.font.size = Pt(20); run.font.name = "STSong"; run._element.rPr.rFonts.set(qn("w:eastAsia"), "STSong")
        meta = doc.add_paragraph(f"制单单位：{company_name}    日期：{datetime.now(UTC).date().isoformat()}")
        meta.runs[0].font.name = "STSong"; meta.runs[0]._element.rPr.rFonts.set(qn("w:eastAsia"), "STSong")
        for line in content.splitlines():
            paragraph = doc.add_paragraph(line)
            paragraph.paragraph_format.space_after = Pt(6)
            if not paragraph.runs:
                paragraph.add_run("")
            paragraph.runs[0].font.name = "STSong"; paragraph.runs[0]._element.rPr.rFonts.set(qn("w:eastAsia"), "STSong")
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_run = footer.add_run(f"{short_name} · 系统生成单据")
        footer_run.font.name = "STSong"; footer_run._element.rPr.rFonts.set(qn("w:eastAsia"), "STSong")
        doc.save(path)

    def _pdf(self, path: Path, profile: CompanyProfileRecord, source_type: str, content: str) -> None:
        company_name = profile.company_name or "示例制造集团"
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        c = canvas.Canvas(str(path), pagesize=A4)
        width, height = A4
        c.setFont("STSong-Light", 12)
        c.drawString(48, height - 48, company_name)
        c.setFont("STSong-Light", 20)
        c.drawCentredString(width / 2, height - 82, {"order": "采购订单", "quotation": "询价单", "contract": "采购合同"}[source_type])
        c.setFont("STSong-Light", 10)
        c.drawString(48, height - 106, f"制单单位：{company_name}    日期：{datetime.now(UTC).date().isoformat()}")
        y = height - 142
        c.setFont("STSong-Light", 11)
        for line in content.splitlines():
            if y < 60:
                c.showPage(); c.setFont("STSong-Light", 11); y = height - 56
            c.drawString(48, y, line[:70]); y -= 20
        c.save()


document_export_service = DocumentExportService()
