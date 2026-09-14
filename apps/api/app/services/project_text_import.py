"""Local extraction only: imported text never executes actions or external requests."""
import hashlib
import io
import re
import zipfile
from datetime import date
from email import policy
from email.parser import BytesParser
from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from fastapi import HTTPException
from lxml.etree import XMLSyntaxError

from app.domain.projects import ProjectImportEvidence, ProjectInput, ProjectTask


class PlainHTML(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []
        self.hidden = 0

    def handle_starttag(self, tag, attrs):
        if tag in ('script', 'style'):
            self.hidden += 1
        if tag in ('br', 'p', 'div', 'li', 'tr'):
            self.parts.append('\n')

    def handle_endtag(self, tag):
        if tag in ('script', 'style') and self.hidden:
            self.hidden -= 1

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


def extract(raw, filename):
    if len(raw) > 5*1024*1024:
        raise HTTPException(422, '文件不能超过5MB')
    extension = Path(filename).suffix.lower()
    notes = []
    try:
        if extension == '.docx':
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if len(archive.infolist()) > 500 or sum(f.file_size for f in archive.infolist()) > 20*1024*1024:
                    raise ValueError('Word解压大小或文件数超限')
            document = Document(io.BytesIO(raw))
            text = '\n'.join([p.text for p in document.paragraphs]+[' | '.join(c.text for c in row.cells) for table in document.tables for row in table.rows])
            notes.append('Word提取正文与表格文本；图片、页眉页脚、批注及修订未识别，请核查原文件。')
        elif extension == '.eml':
            message = BytesParser(policy=policy.default).parsebytes(raw)
            part = message.get_body(preferencelist=('plain', 'html'))
            if part is None:
                raise ValueError('邮件没有可识别正文')
            body = part.get_content()
            if part.get_content_type() == 'text/html':
                parser = PlainHTML()
                parser.feed(body)
                body = ''.join(parser.parts)
            text = f"邮件主题：{message.get('subject', '')}\n发件人：{message.get('from', '')}\n{body}"
            notes.append('仅提取当前邮件正文；附件未自动读取，往来引用内容需人工区分。')
        elif extension == '.txt':
            text = raw.decode('utf-8-sig')
        else:
            raise ValueError('支持UTF-8 TXT、EML邮件和DOCX；扫描件、PDF、DOC和MSG暂不支持')
        if not isinstance(text, str) or not text.strip():
            raise ValueError('文件未提取到文本；扫描图片请先转为文本')
        if len(text) > 20000:
            raise ValueError('提取文本超过20000字，请按项目拆分材料')
        return text, notes
    except HTTPException:
        raise
    except (ValueError, UnicodeDecodeError, zipfile.BadZipFile, KeyError, TypeError, XMLSyntaxError):
        raise HTTPException(422, '文件无法解析或格式不受支持；请使用有效的UTF-8 TXT、EML或DOCX（解压不超过20MB）') from None


def preview(text, filename='粘贴文本', notes=None):
    if not text.strip() or len(text) > 20000:
        raise HTTPException(422, '请输入1—20000字的项目资料')
    fingerprint = hashlib.sha256(text.encode()).hexdigest()
    warnings = list(notes or []) + ['当前使用本地规则提取，不是已训练的专有模型；所有条目均为待确认草案。', '合同价款、付款节点不自动作为成本预算；项目经理、责任人、预算、日历和依赖关系须人工完善。', '日期只提取任务行中的完整年月日；相对日期、多日期或无日期不猜测。']
    named = re.search(r'(?:项目名称|项目名)\s*[:：]\s*([^\n\r]+)', text)
    subject = re.search(r'邮件主题\s*[:：]\s*([^\n\r]+)', text)
    name = (named.group(1) if named else subject.group(1) if subject else '待完善项目')[:200]
    tasks = []
    for number, line in enumerate(text.splitlines(), 1):
        line = line.strip()
        if re.match(r'^(?:项目名称|项目名|邮件主题|发件人|合同金额)\s*[:：]', line):
            continue
        if not line or not re.search(r'交付|验收|里程碑|任务[：:]|完成.*(?:设计|采购|安装|调试)|(?:设计|采购|安装|调试).*完成', line):
            continue
        if len(tasks) >= 100:
            warnings.append('仅提取前100条任务候选，其余内容请人工补充。')
            break
        matches = re.findall(r'(\d{4})[年/-](\d{1,2})[月/-](\d{1,2})日?', line)
        finish = None
        if len(matches) == 1:
            try:
                finish = date(*(int(v) for v in matches[0]))
            except ValueError:
                pass
        tasks.append(ProjectTask(name=line[:200], finish=finish, source=f'{filename} · 第{number}行（候选，待确认）：{line}'[:1000]))
    if not tasks:
        warnings.append('未识别到明确任务；请根据原文补充工作包、任务和验收节点。')
    evidence = ProjectImportEvidence(filename=filename[:200], fingerprint=fingerprint, source_text=text, warnings=warnings)
    # Placeholder code must be replaced by the operator before publishing.
    draft = ProjectInput(code='待填写-'+fingerprint[:8], name=name, tasks=tasks, import_evidence=evidence)
    return {'draft': draft.model_dump(mode='json'), 'warnings': warnings, 'fingerprint': fingerprint}
