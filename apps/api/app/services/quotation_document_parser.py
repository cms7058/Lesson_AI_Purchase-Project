"""Local quotation extraction. Files are treated as data and never executed."""

import csv
import io
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document


HEADERS = {
    "material_code": {"物料编码", "物料号", "料号", "materialcode", "itemcode", "code"},
    "unit_price": {"未税单价", "含税单价", "报价", "单价", "unitprice", "price"},
    "tax_rate": {"税率", "taxrate", "tax"},
    "logistics_cost": {"物流费", "运费", "物流成本", "freight", "logisticscost"},
}


def _normal(value) -> str:
    return re.sub(r"[^0-9a-z\u4e00-\u9fff]+", "", str(value or "").lower())


def _number(value, default=0.0) -> float:
    text = str(value or "").replace(",", "").replace("￥", "").strip()
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    return float(match.group()) if match else default


def _tax(value) -> float:
    number = _number(value, 0.13)
    if "%" in str(value) or number > 1:
        number /= 100
    return min(1.0, max(0.0, number))


def _decode(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030", "utf-16"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _check_office_archive(raw: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            members = archive.infolist()
            if len(members) > 1000 or sum(item.file_size for item in members) > 50 * 1024 * 1024:
                raise ValueError("Office文件展开后过大或文件项过多")
    except zipfile.BadZipFile as exc:
        raise ValueError("Office文件结构无效") from exc


def _delimited_rows(text: str) -> list[list[str]]:
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
        rows = list(csv.reader(io.StringIO(text), dialect))
    except csv.Error:
        rows = [re.split(r"\s{2,}|\t|,", line.strip()) for line in text.splitlines() if line.strip()]
    return [[str(cell).strip() for cell in row] for row in rows if any(str(cell).strip() for cell in row)]


def _docx_rows(raw: bytes) -> tuple[list[list[str]], str]:
    document = Document(io.BytesIO(raw))
    rows = [[cell.text.strip() for cell in row.cells] for table in document.tables for row in table.rows]
    text = "\n".join(p.text for p in document.paragraphs)
    return rows, text


def _xlsx_rows(raw: bytes) -> tuple[list[list[str]], str]:
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        shared = []
        if "xl/sharedStrings.xml" in archive.namelist():
            root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
            shared = ["".join(node.text or "" for node in item.iter() if node.tag.endswith("}t")) for item in root]
        sheet_names = sorted(name for name in archive.namelist() if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"))
        if not sheet_names:
            return [], ""
        root = ET.fromstring(archive.read(sheet_names[0]))
        rows = []
        for row in (node for node in root.iter() if node.tag.endswith("}row")):
            cells = []
            for cell in (node for node in row if node.tag.endswith("}c")):
                value = next((node.text or "" for node in cell.iter() if node.tag.endswith("}v")), "")
                if cell.attrib.get("t") == "s" and value.isdigit() and int(value) < len(shared):
                    value = shared[int(value)]
                elif cell.attrib.get("t") == "inlineStr":
                    value = "".join(node.text or "" for node in cell.iter() if node.tag.endswith("}t"))
                cells.append(value)
            if cells:
                rows.append(cells)
        return rows, "\n".join("\t".join(row) for row in rows)


def _column_map(header: list[str]) -> dict[str, int]:
    result = {}
    normalized = [_normal(cell) for cell in header]
    for field, aliases in HEADERS.items():
        for index, value in enumerate(normalized):
            if value in aliases or any(alias in value for alias in aliases if len(alias) >= 3):
                result[field] = index
                break
    return result


def parse_quotation(raw: bytes, filename: str, rfq_lines: list) -> dict:
    suffix = Path(filename).suffix.lower()
    if suffix in {".csv", ".txt"}:
        text = _decode(raw)
        rows = _delimited_rows(text)
    elif suffix == ".docx":
        _check_office_archive(raw)
        rows, text = _docx_rows(raw)
    elif suffix == ".xlsx":
        _check_office_archive(raw)
        try:
            rows, text = _xlsx_rows(raw)
        except (zipfile.BadZipFile, ET.ParseError) as exc:
            raise ValueError("Excel文件结构无效") from exc
    else:
        raise ValueError("当前自动识别支持 CSV、TXT、DOCX、XLSX；PDF或图片请先转为可复制表格文件")

    header_index = -1
    columns = {}
    for index, row in enumerate(rows[:20]):
        candidate = _column_map(row)
        if "material_code" in candidate and "unit_price" in candidate:
            header_index, columns = index, candidate
            break
    if header_index < 0:
        raise ValueError("未识别到“物料编码”和“单价”列，请使用标准列名后重试")

    found = {}
    for row in rows[header_index + 1:]:
        code_index = columns["material_code"]
        if code_index >= len(row) or not row[code_index].strip():
            continue
        code = _normal(row[code_index])
        found[code] = {
            "unit_price": _number(row[columns["unit_price"]]) if columns["unit_price"] < len(row) else 0,
            "tax_rate": _tax(row[columns["tax_rate"]]) if columns.get("tax_rate", 10**6) < len(row) else 0.13,
            "logistics_cost": _number(row[columns["logistics_cost"]]) if columns.get("logistics_cost", 10**6) < len(row) else 0,
        }

    extracted = []
    missing = []
    for line in rfq_lines:
        values = found.get(_normal(line.material_code))
        if values is None:
            missing.append(line.material_code)
            values = {"unit_price": 0, "tax_rate": 0.13, "logistics_cost": 0}
        extracted.append({"material_code": line.material_code, "material_name": line.material_name, **values, "matched": values["unit_price"] > 0})

    def integer(pattern: str, default: int) -> int:
        match = re.search(pattern, text, re.IGNORECASE)
        return int(match.group(1)) if match else default

    matched = len(extracted) - len(missing)
    confidence = round(matched / max(len(extracted), 1), 2)
    warnings = (["未匹配物料：" + "、".join(missing)] if missing else [])
    if "tax_rate" not in columns:
        warnings.append("未识别税率列，已使用默认税率13%")
    if "logistics_cost" not in columns:
        warnings.append("未识别物流费用列，已按0填入")
    return {
        "filename": Path(filename).name,
        "confidence": confidence,
        "delivery_days": integer(r"(?:交期|delivery)[^0-9]{0,8}(\d{1,3})", 14),
        "validity_days": integer(r"(?:有效期|validity)[^0-9]{0,8}(\d{1,3})", 30),
        "lines": extracted,
        "warnings": warnings,
        "requires_confirmation": True,
    }
