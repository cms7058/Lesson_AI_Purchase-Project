from io import BytesIO
from pathlib import Path
from urllib.parse import quote

from docx import Document
from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile, status
from fastapi.responses import FileResponse, StreamingResponse
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.security import CurrentUser, UserRole, get_current_user
from app.domain.common import Page
from app.domain.templates import (
    BusinessTemplate,
    BusinessTemplateCreate,
    BusinessTemplateUpdate,
    RenderedDocument,
    WordTemplateValidation,
)
from app.services.audit_service import write_audit_log
from app.services.template_service import template_service
from app.services.word_template_service import (
    WordTemplateError,
    create_word_template,
    get_word_file,
    validation_report,
)

router = APIRouter(prefix="/templates", tags=["templates"])


def _require_template_manager(current_user: CurrentUser) -> None:
    if current_user.role not in {UserRole.ADMIN, UserRole.PROCUREMENT_MANAGER}:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="当前角色无权管理业务模板")


@router.get("", response_model=Page[BusinessTemplate])
def list_templates(
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    db: Session = Depends(get_db),
) -> Page[BusinessTemplate]:
    items, total = template_service.list_templates(db, page, page_size)
    return Page(items=list(items), page=page, page_size=page_size, total=total)


@router.post("", response_model=BusinessTemplate, status_code=status.HTTP_201_CREATED)
def create_template(
    payload: BusinessTemplateCreate,
    db: Session = Depends(get_db),
    current_user: CurrentUser = Depends(get_current_user),
) -> BusinessTemplate:
    if current_user.role not in {UserRole.ADMIN, UserRole.PROCUREMENT_MANAGER}:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="当前角色无权管理业务模板")
    template = template_service.create_template(db, payload, current_user.user_id)
    write_audit_log(
        db,
        actor_id=current_user.user_id,
        actor_role=current_user.role,
        action="create",
        resource_type="business_template",
        resource_id=str(template.id),
        detail=f"创建{template.template_type}模板 {template.name} v{template.version}",
    )
    db.commit()
    return template


@router.post("/word-contract", response_model=BusinessTemplate, status_code=status.HTTP_201_CREATED)
async def upload_word_contract_template(
    name: str = Form(min_length=2, max_length=120),
    version: str = Form(default="1.0", min_length=1, max_length=24),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: CurrentUser = Depends(get_current_user),
) -> BusinessTemplate:
    _require_template_manager(current_user)
    try:
        template_record = create_word_template(
            db, name=name, version=version, data=await file.read(),
            filename=file.filename or "template.docx", created_by=current_user.user_id,
        )
    except WordTemplateError as error:
        raise HTTPException(status_code=422, detail=str(error)) from error
    write_audit_log(
        db, actor_id=current_user.user_id, actor_role=current_user.role, action="upload",
        resource_type="business_template", resource_id=template_record.id,
        detail=f"上传Word合同模板 {template_record.name} v{template_record.version}",
    )
    db.commit()
    result = template_service.get_template(db, template_record.id)
    assert result is not None
    return result


@router.get("/word-contract/sample")
def download_word_contract_sample() -> StreamingResponse:
    document = Document()
    document.add_heading("采购合同 {{ contract_no }}", level=0)
    document.add_paragraph("甲方：{{ buyer.company_name }}")
    document.add_paragraph("乙方：{{ supplier.company_name }}")
    document.add_paragraph("项目：{{ project_name }}")
    table = document.add_table(rows=1, cols=7)
    for cell, value in zip(table.rows[0].cells, ["序号", "物料编码", "物料名称", "规格", "数量", "未税单价", "含税金额"], strict=True):
        cell.text = value
    start = table.add_row()
    start.cells[0].text = "{%tr for item in items %}"
    row = table.add_row()
    for cell, value in zip(row.cells, ["{{ item.index }}", "{{ item.material_code }}", "{{ item.material_name }}", "{{ item.specification }}", "{{ item.quantity }} {{ item.unit }}", "{{ item.unit_price }}", "{{ item.line_total }}"], strict=True):
        cell.text = value
    end = table.add_row()
    end.cells[0].text = "{%tr endfor %}"
    document.add_paragraph("合同含税总额：{{ currency }} {{ total_amount }}")
    document.add_paragraph("付款条件：{{ payment_terms }}")
    document.add_paragraph("交货地点：{{ delivery_address }}")
    output = BytesIO()
    document.save(output)
    output.seek(0)
    filename = quote("AI助力_合同Word模板示例.docx")
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
    )


@router.get("/{template_id}/validation", response_model=WordTemplateValidation)
def get_word_template_validation(
    template_id: str,
    db: Session = Depends(get_db),
) -> WordTemplateValidation:
    word_file = get_word_file(db, template_id)
    if word_file is None:
        raise HTTPException(status_code=404, detail="该模板不是Word模板")
    return WordTemplateValidation(template_id=template_id, **validation_report(word_file))


@router.get("/{template_id}/file")
def download_word_template(template_id: str, db: Session = Depends(get_db)) -> FileResponse:
    word_file = get_word_file(db, template_id)
    if word_file is None or not Path(word_file.storage_path).is_file():
        raise HTTPException(status_code=404, detail="未找到Word模板文件")
    return FileResponse(
        word_file.storage_path,
        filename=word_file.original_filename,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


@router.patch("/{template_id}", response_model=BusinessTemplate)
def update_template(
    template_id: str,
    payload: BusinessTemplateUpdate,
    db: Session = Depends(get_db),
    current_user: CurrentUser = Depends(get_current_user),
) -> BusinessTemplate:
    if current_user.role not in {UserRole.ADMIN, UserRole.PROCUREMENT_MANAGER}:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="当前角色无权修改业务模板")
    template = template_service.update_template(db, template_id, payload)
    if template is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="未找到业务模板")
    write_audit_log(db, actor_id=current_user.user_id, actor_role=current_user.role, action="update", resource_type="business_template", resource_id=template_id, detail=f"更新模板 {template.name}")
    db.commit()
    return template


@router.delete("/{template_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_template(
    template_id: str,
    db: Session = Depends(get_db),
    current_user: CurrentUser = Depends(get_current_user),
) -> None:
    if current_user.role not in {UserRole.ADMIN, UserRole.PROCUREMENT_MANAGER}:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="当前角色无权删除业务模板")
    if not template_service.delete_template(db, template_id):
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="未找到业务模板")
    write_audit_log(db, actor_id=current_user.user_id, actor_role=current_user.role, action="delete", resource_type="business_template", resource_id=template_id, detail="删除业务模板")
    db.commit()


@router.post("/{template_id}/render/orders/{order_id}", response_model=RenderedDocument)
def render_order_document(
    template_id: str,
    order_id: str,
    db: Session = Depends(get_db),
    current_user: CurrentUser = Depends(get_current_user),
) -> RenderedDocument:
    if current_user.role not in {UserRole.ADMIN, UserRole.PROCUREMENT_MANAGER, UserRole.BUYER}:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="当前角色无权生成订单文档")
    document = template_service.render_order(db, template_id, order_id)
    if document is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="未找到模板或采购订单")
    write_audit_log(
        db,
        actor_id=current_user.user_id,
        actor_role=current_user.role,
        action="render",
        resource_type="business_template",
        resource_id=template_id,
        detail=f"依据订单 {order_id} 生成文档预览",
    )
    db.commit()
    return document
